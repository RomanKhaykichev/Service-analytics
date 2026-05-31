"""Build import-compatible XLSX reports from Uzum Seller API."""

from __future__ import annotations

import io
import logging
import time
from dataclasses import dataclass, field
from typing import Any, Callable, Optional

import requests
from openpyxl import Workbook
from openpyxl.styles import Font

from app.services.uzum_time import (
    calendar_date_to_epoch_ms,
    format_datetime,
    last_n_days_range_ms,
    normalize_epoch_ms,
    parse_to_epoch_ms,
    timezone_metadata,
)
from app.utils.column_mappings import (
    CANONICAL_EXPENSES,
    CANONICAL_LEFTOUT_API,
    CANONICAL_SALES,
    CANONICAL_STORAGE,
    EXPENSES_COLUMN_ORDER,
    LEFTOUT_API_COLUMN_ORDER,
    SALES_COLUMN_ORDER,
    STORAGE_COLUMN_ORDER,
)

logger = logging.getLogger(__name__)

API_BASE_URL = "https://api-seller.uzum.uz/api/seller-openapi"
REQUEST_TIMEOUT = 45
PAGE_SIZE = 100
FBS_ORDERS_PAGE_SIZE = 50  # OpenAPI max for GET /v2/fbs/orders
PRODUCT_PAGE_SIZE = 100
MAX_PAGES = 200
# Uzum rate-limits seller API; space out calls and retry on 429.
MIN_REQUEST_INTERVAL_SEC = 0.5
MAX_RETRIES_RATE_LIMIT = 6
MAX_BACKOFF_SEC = 45.0
PAUSE_BETWEEN_SHOPS_SEC = 0.75
SALES_LOOKBACK_DAYS = 15

FINANCE_ORDER_STATUSES = [
    "TO_WITHDRAW",
    "PROCESSING",
    "CANCELED",
    "PARTIALLY_CANCELLED",
]

SALES_STATUS_RU = {
    "TO_WITHDRAW": "Завершен",
    "PROCESSING": "В обработке",
    "CANCELED": "Отменен",
    "PARTIALLY_CANCELLED": "Частично отменен",
}

EXPENSE_SOURCE_RU: dict[str, str] = {
    "STORAGE": "Склад",
    "WAREHOUSE": "Склад",
    "SKLAD": "Склад",
    "OMBOR": "Склад",
    "MARKETING": "Маркетинг",
    "ADVERTISING": "Маркетинг",
    "REKLAMA": "Маркетинг",
    "ADS": "Маркетинг",
}

EXPENSE_STATUS_RU = {
    "CREATED": "Создан",
    "REFUNDED": "Возвращен",
    "CONFIRMED": "Оплачено",
    "CANCELED": "Отменен",
}

EXPENSE_TYPE_RU = {
    "OUTCOME": "Оплата",
    "INCOME": "Возврат",
}

SHEET_NAMES = {
    "sales": "Отчет по продажам",
    "expenses": "Отчет по услугам",
    "storage": "Отчет по хранению",
    "inventory_old": "Отчет по остаткам",
}

FILE_NAMES = {
    "sales": "sells-report_uzum-api.xlsx",
    "expenses": "expenses-report_uzum-api.xlsx",
    "storage": "seller-storage-report_uzum-api.xlsx",
    "inventory_old": "left-out-report_uzum-api.xlsx",
}


@dataclass(frozen=True)
class ExportColumn:
    name: str
    mapped: bool


class ShopUnavailableError(RuntimeError):
    """Uzum rejected a shopId (forbidden-001 / Shop is not available)."""


class UzumRateLimitError(RuntimeError):
    """Uzum returned HTTP 429 (too many requests)."""


def _errors_indicate_unavailable_shop(errors: Any) -> bool:
    if not isinstance(errors, list):
        return False
    for entry in errors:
        if not isinstance(entry, dict):
            continue
        code = str(entry.get("code") or "")
        message = str(entry.get("message") or "")
        if code == "forbidden-001" or "not available" in message.lower():
            return True
    return False


def _detail_is_shop_unavailable(detail: Any) -> bool:
    if _errors_indicate_unavailable_shop(detail):
        return True
    if isinstance(detail, dict):
        return _errors_indicate_unavailable_shop(detail.get("errors"))
    return False


def _format_uzum_error(status: int, detail: Any) -> str:
    if isinstance(detail, dict):
        errors = detail.get("errors")
        if errors:
            return f"Uzum API {status}: {errors}"
        if detail.get("message"):
            return f"Uzum API {status}: {detail['message']}"
        if detail.get("error"):
            return f"Uzum API {status}: {detail['error']}"
    if isinstance(detail, list):
        return f"Uzum API {status}: {detail}"
    return f"Uzum API {status}: {detail}"


def _check_envelope_errors(data: Any) -> None:
    if not isinstance(data, dict):
        return
    errors = data.get("errors")
    if errors:
        if _errors_indicate_unavailable_shop(errors):
            raise ShopUnavailableError(f"Uzum API: {errors}")
        raise RuntimeError(f"Uzum API: {errors}")


def _flatten_order_items(batch: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Unwrap ProductGroupedSellerItem (group=true) into flat line items."""
    flat: list[dict[str, Any]] = []
    for item in batch:
        nested = item.get("items")
        if isinstance(nested, list) and nested:
            parent_title = item.get("productTitle") or ""
            parent_shop = item.get("shopId")
            for sub in nested:
                if not isinstance(sub, dict):
                    continue
                row = dict(sub)
                if not row.get("productTitle"):
                    row["productTitle"] = parent_title
                if row.get("shopId") is None:
                    row["shopId"] = parent_shop
                flat.append(row)
        else:
            flat.append(item)
    return flat


def _shop_label(shop_names: dict[int, str], shop_id: Any) -> str:
    if shop_id is None:
        return ""
    try:
        return shop_names.get(int(shop_id), "")
    except (TypeError, ValueError):
        return ""


def _order_item_timestamp_ms(item: dict[str, Any]) -> Optional[int]:
    for key in ("date", "dateIssued"):
        ts = normalize_epoch_ms(item.get(key))
        if ts is not None:
            return ts
    return None


def _payment_timestamp_ms(payment: dict[str, Any]) -> Optional[int]:
    for key in ("dateService", "dateCreated"):
        ts = parse_to_epoch_ms(payment.get(key))
        if ts is not None:
            return ts
    return None


def _in_date_range(
    ts_ms: Optional[int],
    date_from_ms: Optional[int],
    date_to_ms: Optional[int],
) -> bool:
    if date_from_ms is None and date_to_ms is None:
        return True
    if ts_ms is None:
        return False
    if date_from_ms is not None and ts_ms < date_from_ms:
        return False
    if date_to_ms is not None and ts_ms > date_to_ms:
        return False
    return True


def _filter_by_date_range(
    records: list[dict[str, Any]],
    date_from_ms: Optional[int],
    date_to_ms: Optional[int],
    ts_getter: Callable[[dict[str, Any]], Optional[int]],
) -> list[dict[str, Any]]:
    if date_from_ms is None and date_to_ms is None:
        return records
    return [
        record
        for record in records
        if _in_date_range(ts_getter(record), date_from_ms, date_to_ms)
    ]


def _apply_date_query_params(
    params: dict[str, Any],
    date_from_ms: Optional[int],
    date_to_ms: Optional[int],
    *,
    unit_ms: bool = True,
) -> None:
    if date_from_ms is not None:
        params["dateFrom"] = date_from_ms if unit_ms else date_from_ms // 1000
    if date_to_ms is not None:
        params["dateTo"] = date_to_ms if unit_ms else date_to_ms // 1000


def _format_barcode(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _barcode_keys(barcode: str) -> list[str]:
    if not barcode:
        return []
    keys = [barcode]
    stripped = barcode.lstrip("0")
    if stripped and stripped not in keys:
        keys.append(stripped)
    return keys


# OpenAPI DimensionalGroup.group + типичные коды в выгрузках UZ (KGT/MGT/BGT) -> RU кабинета
_DIMENSIONAL_GROUP_SHORT: dict[str, str] = {
    "СГТ": "S",
    "МГТ": "M",
    "БГТ": "L",
    "SMALL": "S",
    "MEDIUM": "M",
    "LARGE": "L",
    "S": "S",
    "M": "M",
    "L": "L",
}

_DIMENSIONAL_GROUP_TO_RU: dict[str, str] = {
    "SMALL": "СГТ",
    "MEDIUM": "МГТ",
    "LARGE": "БГТ",
    "UNKNOWN": "",
    "KGT": "МГТ",
    "MGT": "МГТ",
    "BGT": "БГТ",
    "SGT": "СГТ",
    "СГТ": "СГТ",
    "МГТ": "МГТ",
    "БГТ": "БГТ",
    "KICHIK": "СГТ",
    "ORTA": "МГТ",
    "KATTA": "БГТ",
}


def _format_dimensional_group(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return ""
        mapped = _DIMENSIONAL_GROUP_TO_RU.get(s.upper())
        return mapped if mapped is not None else s
    if isinstance(value, dict):
        grp = value.get("group")
        if grp is not None:
            mapped = _DIMENSIONAL_GROUP_TO_RU.get(str(grp).upper())
            if mapped is not None:
                return mapped
            return str(grp).strip()
        for key in ("title", "name", "label", "code"):
            part = value.get(key)
            if part:
                return _format_dimensional_group(str(part))
        return ""
    return str(value).strip()


def _dimensional_group_short(label: str) -> str:
    if not label:
        return ""
    return _DIMENSIONAL_GROUP_SHORT.get(label.strip().upper(), label)


def _sku_dimensional_group(sku: dict[str, Any]) -> str:
    for key in (
        "actualDimensionalGroup",
        "dimensionalGroup",
        "paidStorageDimensionalGroup",
    ):
        label = _format_dimensional_group(sku.get(key))
        if label:
            return label
    return ""


def _storage_type_label(cat: Optional["SkuCatalogEntry"]) -> str:
    """«Хранение»: Бесплатное / Платное (как в seller-storage-report кабинета)."""
    if not cat:
        return ""
    if cat.pstorage is True:
        return "Платное"
    if cat.pstorage is False:
        return "Бесплатное"
    amount = _safe_float(cat.paid_storage_amount)
    if amount is not None and amount > 0:
        return "Платное"
    return "Бесплатное"


def _storage_price_per_unit(cat: Optional["SkuCatalogEntry"]) -> Any:
    """Тариф за 1 ед. в день (paidStoragePriceItem) — в кабинете есть и при бесплатном хранении."""
    if not cat or cat.paid_storage_price_item is None:
        return ""
    return cat.paid_storage_price_item


def _storage_fee_1_day(cat: Optional["SkuCatalogEntry"], fbo_stock: Any) -> Any:
    """Всего за 1 день = тариф × остаток FBO (как в seller-storage-report кабинета)."""
    if not cat:
        return ""
    if cat.pstorage is False:
        return 0
    price = _safe_float(cat.paid_storage_price_item)
    stock = _safe_float(fbo_stock if fbo_stock not in (None, "") else cat.quantity_active)
    if price is not None and stock is not None:
        return int(round(price * stock))
    return 0


def _storage_fee_30d(cat: Optional["SkuCatalogEntry"], fbo_stock: Any) -> Any:
    """Всего за 30 дней = (тариф × остаток FBO) × 30, как в seller-storage-report."""
    if not cat:
        return ""
    daily = _safe_float(_storage_fee_1_day(cat, fbo_stock))
    if daily is not None and daily > 0:
        return int(round(daily * 30))
    amount = _safe_float(cat.paid_storage_amount)
    if amount is not None and amount > 0:
        return int(round(amount))
    return 0


def _include_in_storage_report(entry: SkuCatalogEntry) -> bool:
    """Строки отчёта по FBO-хранению — из каталога, не из FBS-остатков."""
    if (_safe_float(entry.quantity_active) or 0) > 0:
        return True
    if (_safe_float(entry.avg_daily_stock) or 0) > 0:
        return True
    if entry.paid_storage_price_item is not None:
        return True
    if entry.pstorage is not None:
        return True
    if entry.dimensional_group:
        return True
    return False


def _storage_turnover_days(avg_stock_15d: Any, avg_sales_15d: Any) -> Any:
    """Оборачиваемость = среднесуточные остатки FBO за 15 дн. / среднесуточные продажи FBO за 15 дн."""
    stock = _safe_float(avg_stock_15d)
    avg = _safe_float(avg_sales_15d)
    if stock is None or avg is None or avg <= 0:
        return ""
    days = stock / avg
    if days >= 100:
        return int(round(days))
    return int(round(days)) if abs(days - round(days)) < 0.05 else round(days, 1)


def _safe_float(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _compute_turnover_days(stock_qty: Any, avg_daily_sales: Any) -> Any:
    """Остаток / среднесуточные продажи (дней), как в left-out-report."""
    stock = _safe_float(stock_qty)
    avg = _safe_float(avg_daily_sales)
    if stock is None or avg is None or avg <= 0:
        return ""
    days = stock / avg
    if days >= 100:
        return int(round(days))
    return round(days, 1)


PRODUCT_STATUS_RU: dict[str, str] = {
    "IN_STOCK": "В продаже",
    "READY_TO_SEND": "Готов к отправке",
    "SENT": "Отправлен",
    "RUN_OUT": "Закончился",
    "BLOCKED": "Заблокирован",
    "SKU_BLOCKED": "SKU заблокирован",
    "ARCHIVED": "В архиве",
    "DELETED": "Удалён",
    "NO_SKU": "Нет SKU",
    "NOT_READY_TO_SEND": "Не готов к отправке",
    "PERM_BANNED": "Заблокирован навсегда",
}


def _product_status_label(product: dict[str, Any]) -> str:
    status = product.get("status")
    if isinstance(status, dict):
        title = status.get("title")
        if title:
            return str(title)
        value = status.get("value")
        if value:
            return PRODUCT_STATUS_RU.get(str(value), str(value))
    return ""


def _qty_int(value: Any) -> int:
    parsed = _safe_float(value)
    return int(parsed) if parsed is not None else 0


def _optional_number(value: Any) -> Any:
    if value in (None, ""):
        return None
    return value


def _invoice_is_pending(invoice: dict[str, Any]) -> bool:
    """Накладная FBO, по которой ещё есть товар к отправке на склад."""
    if invoice.get("dateAccepted"):
        return False
    status = invoice.get("invoiceStatus") or {}
    value = str(status.get("value") or invoice.get("status") or "").upper()
    if value in ("ACCEPTED", "COMPLETED", "CANCELLED", "CANCELED", "REJECTED", "DECLINED"):
        return False
    total_to = _qty_int(invoice.get("totalToStock"))
    total_acc = _qty_int(invoice.get("totalAccepted"))
    if total_to > 0 and total_acc >= total_to and invoice.get("dateAccepted"):
        return False
    return True


def _sku_label_match(label: str, entry: "SkuCatalogEntry") -> bool:
    if not label:
        return False
    for candidate in (
        entry.article,
        entry.seller_item_code,
        entry.sku_title,
        entry.sku_full_title,
    ):
        if not candidate:
            continue
        if label == candidate or label in candidate or candidate in label:
            return True
    return False


def _format_sales_rate(value: float) -> Any:
    if value <= 0:
        return ""
    if abs(value - round(value)) < 0.01:
        return int(round(value))
    return round(value, 2)


@dataclass
class SkuCatalogEntry:
    shop_id: int
    sku_id: Optional[int]
    product_id: Optional[int]
    barcode: str
    category: str
    dimensional_group: str
    purchase_price: Any
    sell_price: Any
    paid_storage_price_item: Any
    paid_storage_amount: Any
    avg_daily_sales: Any
    avg_daily_stock: Any
    quantity_active: Any
    quantity_pending: Any
    quantity_fbs: Any
    quantity_additional: Any
    pstorage: Optional[bool]
    sku_title: str
    sku_full_title: str
    article: str
    seller_item_code: str
    product_title: str
    quantity_defected: Any
    quantity_returned: Any
    characteristics: str
    commission: Any
    product_status: str
    shop_name: str
    preview_image: str
    quantity_on_photo_studio: Any


@dataclass
class ProductCatalogIndex:
    by_sku_id: dict[tuple[int, int], SkuCatalogEntry] = field(default_factory=dict)
    by_sku_id_global: dict[int, SkuCatalogEntry] = field(default_factory=dict)
    by_product_id: dict[tuple[int, int], SkuCatalogEntry] = field(default_factory=dict)
    by_barcode: dict[str, list[SkuCatalogEntry]] = field(default_factory=dict)
    skus_by_product: dict[tuple[int, int], list[SkuCatalogEntry]] = field(default_factory=dict)

    @classmethod
    def empty(cls) -> ProductCatalogIndex:
        return cls()

    @property
    def size(self) -> int:
        return len(self.by_sku_id_global)

    def enrich_storage_fields(self) -> None:
        """Дозаполнить габарит и тариф из других SKU того же товара / магазина."""
        for skus in self.skus_by_product.values():
            ref_dg = ""
            ref_price: Any = None
            for sku in skus:
                if sku.dimensional_group:
                    ref_dg = sku.dimensional_group
                    break
            for sku in skus:
                if sku.paid_storage_price_item is not None:
                    ref_price = sku.paid_storage_price_item
                    break
            for sku in skus:
                if not sku.dimensional_group and ref_dg:
                    sku.dimensional_group = ref_dg
                if sku.paid_storage_price_item is None and ref_price is not None:
                    sku.paid_storage_price_item = ref_price

        dg_price_by_shop: dict[tuple[int, str], Any] = {}
        for entry in self.by_sku_id.values():
            if entry.dimensional_group and entry.paid_storage_price_item is not None:
                key = (entry.shop_id, entry.dimensional_group)
                dg_price_by_shop.setdefault(key, entry.paid_storage_price_item)

        for entry in self.by_sku_id.values():
            if entry.paid_storage_price_item is None and entry.dimensional_group:
                fallback = dg_price_by_shop.get((entry.shop_id, entry.dimensional_group))
                if fallback is not None:
                    entry.paid_storage_price_item = fallback

    def enrich_purchase_prices_from_orders(self, order_items: list[dict[str, Any]]) -> int:
        """Себестоимость из финансовых продаж (как в sells-report), если в каталоге пусто."""
        filled = 0
        for item in order_items:
            if not isinstance(item, dict):
                continue
            price = item.get("purchasePrice")
            if price in (None, ""):
                continue
            entry = self.lookup_order_item(item)
            if entry is None or entry.purchase_price is not None:
                continue
            entry.purchase_price = price
            filled += 1
        return filled

    def enrich_from_fbo_invoices(self, invoices: list[dict[str, Any]]) -> tuple[int, int]:
        """К отправке и себестоимость из открытых FBO-накладных (/v1/invoice)."""
        to_ship_filled = 0
        price_filled = 0
        pending_by_sku: dict[tuple[int, int], int] = {}

        for invoice in invoices:
            if not isinstance(invoice, dict) or not _invoice_is_pending(invoice):
                continue
            shop_raw = invoice.get("shopId")
            if shop_raw is None:
                continue
            try:
                shop_id = int(shop_raw)
            except (TypeError, ValueError):
                continue

            products = invoice.get("productForInvoiceDto") or []
            if not isinstance(products, list):
                continue
            for product in products:
                if not isinstance(product, dict):
                    continue
                sku_list = product.get("skuForInvoiceDtoList") or []
                if not isinstance(sku_list, list):
                    continue
                product_price = _optional_number(product.get("purchasePrice"))
                for sku in sku_list:
                    if not isinstance(sku, dict):
                        continue
                    entry = self._match_invoice_sku(shop_id, sku, product)
                    if entry is None or entry.sku_id is None:
                        continue
                    sku_price = _optional_number(sku.get("purchasePrice")) or product_price
                    if entry.purchase_price is None and sku_price is not None:
                        entry.purchase_price = sku_price
                        price_filled += 1
                    qty = _qty_int(sku.get("quantityToStock"))
                    if qty <= 0:
                        continue
                    key = (shop_id, entry.sku_id)
                    pending_by_sku[key] = pending_by_sku.get(key, 0) + qty

        for (shop_id, sku_id), qty in pending_by_sku.items():
            entry = self.by_sku_id.get((shop_id, sku_id))
            if entry is None:
                continue
            if _qty_int(entry.quantity_pending) > 0:
                continue
            entry.quantity_pending = qty
            to_ship_filled += 1
        return to_ship_filled, price_filled

    def _match_invoice_sku(
        self,
        shop_id: int,
        sku: dict[str, Any],
        product: dict[str, Any],
    ) -> Optional[SkuCatalogEntry]:
        raw_id = sku.get("id")
        if raw_id is not None:
            try:
                entry = self.by_sku_id.get((shop_id, int(raw_id)))
                if entry:
                    return entry
            except (TypeError, ValueError):
                pass

        title = str(sku.get("skuTitle") or product.get("skuTitle") or "")
        if title:
            for entry in self.by_sku_id.values():
                if entry.shop_id == shop_id and _sku_label_match(title, entry):
                    return entry
        return None

    def enrich_purchase_prices_from_product_skus(self) -> None:
        for skus in self.skus_by_product.values():
            ref: Any = None
            for sku in skus:
                if sku.purchase_price is not None:
                    ref = sku.purchase_price
                    break
            if ref is None:
                continue
            for sku in skus:
                if sku.purchase_price is None:
                    sku.purchase_price = ref

    def add_shop_products(self, shop_id: int, products: list[dict[str, Any]]) -> None:
        for product in products:
            if not isinstance(product, dict):
                continue
            category = str(product.get("category") or "")
            product_id_raw = product.get("productId")
            product_id: Optional[int] = None
            if product_id_raw is not None:
                try:
                    product_id = int(product_id_raw)
                except (TypeError, ValueError):
                    product_id = None
            product_title = str(product.get("title") or product.get("skuTitle") or "")
            product_status = _product_status_label(product)
            sku_list = product.get("skuList") or []
            if not isinstance(sku_list, list):
                continue
            for sku in sku_list:
                if not isinstance(sku, dict):
                    continue
                sku_id_raw = sku.get("skuId")
                sku_id: Optional[int] = None
                if sku_id_raw is not None:
                    try:
                        sku_id = int(sku_id_raw)
                    except (TypeError, ValueError):
                        sku_id = None
                barcode = _format_barcode(sku.get("barcode"))
                entry = SkuCatalogEntry(
                    shop_id=shop_id,
                    sku_id=sku_id,
                    product_id=product_id,
                    barcode=barcode,
                    category=category,
                    dimensional_group=_sku_dimensional_group(sku),
                    purchase_price=_optional_number(sku.get("purchasePrice")),
                    sell_price=sku.get("price"),
                    paid_storage_price_item=sku.get("paidStoragePriceItem"),
                    paid_storage_amount=sku.get("paidStorageAmount"),
                    avg_daily_sales=sku.get("avgdsales"),
                    avg_daily_stock=sku.get("avgdquantity"),
                    quantity_active=sku.get("quantityActive"),
                    quantity_pending=sku.get("quantityPending"),
                    quantity_fbs=sku.get("quantityFbs"),
                    quantity_additional=sku.get("quantityAdditional"),
                    pstorage=sku.get("pstorage"),
                    sku_title=str(sku.get("skuTitle") or ""),
                    sku_full_title=str(sku.get("skuFullTitle") or ""),
                    article=str(sku.get("article") or ""),
                    seller_item_code=str(sku.get("sellerItemCode") or ""),
                    product_title=str(sku.get("productTitle") or product_title),
                    quantity_defected=sku.get("quantityDefected"),
                    quantity_returned=sku.get("quantityReturned"),
                    characteristics=str(sku.get("characteristics") or "").strip(),
                    commission=sku.get("commission"),
                    product_status=product_status,
                    shop_name="",
                    preview_image=str(sku.get("previewImage") or ""),
                    quantity_on_photo_studio=sku.get("quantityOnPhotoStudio"),
                )
                if sku_id is not None:
                    self.by_sku_id[(shop_id, sku_id)] = entry
                    self.by_sku_id_global[sku_id] = entry
                if product_id is not None:
                    key = (shop_id, product_id)
                    self.by_product_id.setdefault(key, entry)
                    self.skus_by_product.setdefault(key, []).append(entry)
                for key in _barcode_keys(barcode):
                    self.by_barcode.setdefault(key, []).append(entry)

    def lookup_order_item(self, item: dict[str, Any]) -> Optional[SkuCatalogEntry]:
        shop_raw = item.get("shopId")
        if shop_raw is None:
            return self._lookup_without_shop(item)
        try:
            shop_id = int(shop_raw)
        except (TypeError, ValueError):
            return self._lookup_without_shop(item)

        sku_raw = item.get("skuId")
        if sku_raw is not None:
            try:
                entry = self.by_sku_id.get((shop_id, int(sku_raw)))
                if entry:
                    return entry
            except (TypeError, ValueError):
                pass

        product_raw = item.get("productId")
        product_id: Optional[int] = None
        if product_raw is not None:
            try:
                product_id = int(product_raw)
            except (TypeError, ValueError):
                product_id = None

        if product_id is not None:
            key = (shop_id, product_id)
            labels = [
                str(item.get("skuTitle") or ""),
                str(item.get("skuCharTitle") or ""),
                str(item.get("skuCharValue") or ""),
            ]
            for candidate in self.skus_by_product.get(key, []):
                for label in labels:
                    if _sku_label_match(label, candidate):
                        return candidate
            entry = self.by_product_id.get(key)
            if entry:
                return entry

        return self._lookup_without_shop(item)

    def _lookup_without_shop(self, item: dict[str, Any]) -> Optional[SkuCatalogEntry]:
        sku_raw = item.get("skuId")
        if sku_raw is not None:
            try:
                entry = self.by_sku_id_global.get(int(sku_raw))
                if entry:
                    return entry
            except (TypeError, ValueError):
                pass
        for key in _barcode_keys(_format_barcode(item.get("barcode"))):
            entries = self.by_barcode.get(key)
            if entries:
                return entries[0]
        return None

    def lookup_stock_row(self, stock: dict[str, Any]) -> Optional[SkuCatalogEntry]:
        sku_raw = stock.get("skuId")
        if sku_raw is not None:
            try:
                entry = self.by_sku_id_global.get(int(sku_raw))
                if entry:
                    return entry
            except (TypeError, ValueError):
                pass
        for key in _barcode_keys(_format_barcode(stock.get("barcode"))):
            entries = self.by_barcode.get(key)
            if entries:
                return entries[0]
        return None


class UzumApiClient:
    def __init__(self, api_key: str):
        self._headers = {
            "Authorization": api_key,
            "Accept": "application/json",
            "Accept-Language": "ru-RU",
        }
        self.warnings: list[str] = []
        self._last_request_at: float = 0.0
        self._shops_cache: Optional[list[dict[str, Any]]] = None

    def _throttle(self) -> None:
        elapsed = time.monotonic() - self._last_request_at
        if elapsed < MIN_REQUEST_INTERVAL_SEC:
            time.sleep(MIN_REQUEST_INTERVAL_SEC - elapsed)

    def _retry_after_seconds(self, response: requests.Response, attempt: int) -> float:
        raw = response.headers.get("Retry-After") or response.headers.get("retry-after")
        if raw:
            try:
                return min(float(raw), MAX_BACKOFF_SEC)
            except ValueError:
                pass
        return min(2.0 * (2**attempt), MAX_BACKOFF_SEC)

    def get(self, path: str, params: Optional[dict[str, Any]] = None) -> Any:
        url = f"{API_BASE_URL.rstrip('/')}{path if path.startswith('/') else '/' + path}"
        last_error: Optional[RuntimeError] = None

        for attempt in range(MAX_RETRIES_RATE_LIMIT + 1):
            self._throttle()
            response = requests.get(
                url,
                headers=self._headers,
                params=params,
                timeout=REQUEST_TIMEOUT,
            )
            self._last_request_at = time.monotonic()

            if response.status_code == 429:
                wait_sec = self._retry_after_seconds(response, attempt)
                logger.warning(
                    "Uzum API 429 on %s, retry %s/%s after %.1fs",
                    path,
                    attempt + 1,
                    MAX_RETRIES_RATE_LIMIT,
                    wait_sec,
                )
                if attempt >= MAX_RETRIES_RATE_LIMIT:
                    raise UzumRateLimitError(
                        "Uzum API временно ограничил частоту запросов (429). "
                        "Подождите 1–2 минуты и повторите выгрузку."
                    )
                time.sleep(wait_sec)
                continue

            if response.status_code >= 400:
                try:
                    detail = response.json()
                except ValueError:
                    detail = response.text
                if _detail_is_shop_unavailable(detail):
                    raise ShopUnavailableError(_format_uzum_error(response.status_code, detail))
                last_error = RuntimeError(_format_uzum_error(response.status_code, detail))
                break

            if not response.content:
                return None
            data = response.json()
            _check_envelope_errors(data)
            return data

        if last_error:
            raise last_error
        raise UzumRateLimitError(
            "Uzum API временно ограничил частоту запросов (429). "
            "Подождите 1–2 минуты и повторите выгрузку."
        )

    def list_shops(self) -> list[dict[str, Any]]:
        if self._shops_cache is not None:
            return self._shops_cache
        data = self.get("/v1/shops")
        if isinstance(data, list):
            self._shops_cache = data
            return data
        self._shops_cache = []
        return []

    def list_shop_ids(self) -> list[int]:
        ids: list[int] = []
        for shop in self.list_shops():
            raw = shop.get("id")
            if raw is None:
                continue
            try:
                ids.append(int(raw))
            except (TypeError, ValueError):
                continue
        return ids

    def resolve_shop_ids(self, shop_ids: Optional[list[int]] = None) -> list[int]:
        if shop_ids:
            return shop_ids
        resolved = self.list_shop_ids()
        if not resolved:
            raise RuntimeError(
                "Не найдены магазины в Uzum API (/v1/shops). Проверьте API-ключ и права доступа."
            )
        return resolved

    def _fetch_orders_for_shop(
        self,
        shop_id: int,
        date_from_ms: Optional[int],
        date_to_ms: Optional[int],
        *,
        unit_ms: bool = True,
    ) -> list[dict[str, Any]]:
        items: list[dict[str, Any]] = []
        page = 0
        while page < MAX_PAGES:
            params: dict[str, Any] = {
                "page": page,
                "size": PAGE_SIZE,
                "group": False,
                "shopIds": [shop_id],
                "statuses": FINANCE_ORDER_STATUSES,
            }
            _apply_date_query_params(params, date_from_ms, date_to_ms, unit_ms=unit_ms)
            data = self.get("/v1/finance/orders", params)
            raw_batch = (data or {}).get("orderItems") or []
            batch = _flatten_order_items(raw_batch)
            if not batch:
                break
            items.extend(batch)
            total = (data or {}).get("totalElements")
            if total is not None and len(items) >= int(total):
                break
            if len(raw_batch) < PAGE_SIZE:
                break
            page += 1
        return items

    def _fetch_expenses_for_shop(
        self,
        shop_id: int,
        date_from_ms: Optional[int],
        date_to_ms: Optional[int],
        *,
        unit_ms: bool = True,
    ) -> list[dict[str, Any]]:
        payments: list[dict[str, Any]] = []
        page = 0
        while page < MAX_PAGES:
            params: dict[str, Any] = {
                "page": page,
                "size": PAGE_SIZE,
                "shopIds": [shop_id],
            }
            _apply_date_query_params(params, date_from_ms, date_to_ms, unit_ms=unit_ms)
            data = self.get("/v1/finance/expenses", params)
            payload = (data or {}).get("payload") or data or {}
            if isinstance(payload, dict):
                errors = payload.get("errors")
                if errors:
                    if _errors_indicate_unavailable_shop(errors):
                        raise ShopUnavailableError(f"Uzum API: {errors}")
                    raise RuntimeError(f"Uzum API: {errors}")
            batch: list[dict[str, Any]] = []
            if isinstance(payload, dict):
                batch = payload.get("payments") or []
            if not batch:
                break
            payments.extend(batch)
            if len(batch) < PAGE_SIZE:
                break
            page += 1
        return payments

    def _fetch_orders_bulk(
        self,
        shop_ids: list[int],
        date_from_ms: Optional[int],
        date_to_ms: Optional[int],
        *,
        unit_ms: bool = True,
    ) -> list[dict[str, Any]]:
        """One paginated stream for all available shops (fewer HTTP calls)."""
        items: list[dict[str, Any]] = []
        page = 0
        while page < MAX_PAGES:
            params: dict[str, Any] = {
                "page": page,
                "size": PAGE_SIZE,
                "group": False,
                "shopIds": shop_ids,
                "statuses": FINANCE_ORDER_STATUSES,
            }
            _apply_date_query_params(params, date_from_ms, date_to_ms, unit_ms=unit_ms)
            data = self.get("/v1/finance/orders", params)
            raw_batch = (data or {}).get("orderItems") or []
            batch = _flatten_order_items(raw_batch)
            if not batch:
                break
            items.extend(batch)
            total = (data or {}).get("totalElements")
            if total is not None and len(items) >= int(total):
                break
            if len(raw_batch) < PAGE_SIZE:
                break
            page += 1
        return items

    def _fetch_expenses_bulk(
        self,
        shop_ids: list[int],
        date_from_ms: Optional[int],
        date_to_ms: Optional[int],
        *,
        unit_ms: bool = True,
    ) -> list[dict[str, Any]]:
        payments: list[dict[str, Any]] = []
        page = 0
        while page < MAX_PAGES:
            params: dict[str, Any] = {
                "page": page,
                "size": PAGE_SIZE,
                "shopIds": shop_ids,
            }
            _apply_date_query_params(params, date_from_ms, date_to_ms, unit_ms=unit_ms)
            data = self.get("/v1/finance/expenses", params)
            payload = (data or {}).get("payload") or data or {}
            if isinstance(payload, dict):
                errors = payload.get("errors")
                if errors:
                    if _errors_indicate_unavailable_shop(errors):
                        raise ShopUnavailableError(f"Uzum API: {errors}")
                    raise RuntimeError(f"Uzum API: {errors}")
            batch: list[dict[str, Any]] = []
            if isinstance(payload, dict):
                batch = payload.get("payments") or []
            if not batch:
                break
            payments.extend(batch)
            if len(batch) < PAGE_SIZE:
                break
            page += 1
        return payments

    def _probe_shop_for_orders(self, shop_id: int) -> bool:
        params: dict[str, Any] = {
            "page": 0,
            "size": 1,
            "group": False,
            "shopIds": [shop_id],
            "statuses": FINANCE_ORDER_STATUSES,
        }
        try:
            self.get("/v1/finance/orders", params)
            return True
        except ShopUnavailableError:
            return False

    def _probe_shop_for_expenses(self, shop_id: int) -> bool:
        params: dict[str, Any] = {"page": 0, "size": 1, "shopIds": [shop_id]}
        try:
            self.get("/v1/finance/expenses", params)
            return True
        except ShopUnavailableError:
            return False

    def _fetch_orders_for_shops(
        self,
        shop_ids: list[int],
        names: dict[int, str],
        date_from_ms: Optional[int],
        date_to_ms: Optional[int],
        *,
        unit_ms: bool = True,
    ) -> list[dict[str, Any]]:
        if len(shop_ids) == 1:
            return self._fetch_orders_for_shop(
                shop_ids[0], date_from_ms, date_to_ms, unit_ms=unit_ms
            )
        try:
            return self._fetch_orders_bulk(
                shop_ids, date_from_ms, date_to_ms, unit_ms=unit_ms
            )
        except ShopUnavailableError:
            self.warnings.append(
                "Сводная выгрузка недоступна, загружаем магазины по одному (медленнее)."
            )
            return self._collect_per_shop(
                shop_ids,
                names,
                lambda sid: self._fetch_orders_for_shop(
                    sid, date_from_ms, date_to_ms, unit_ms=unit_ms
                ),
            )

    def _fetch_expenses_for_shops(
        self,
        shop_ids: list[int],
        names: dict[int, str],
        date_from_ms: Optional[int],
        date_to_ms: Optional[int],
        *,
        unit_ms: bool = True,
    ) -> list[dict[str, Any]]:
        if len(shop_ids) == 1:
            return self._fetch_expenses_for_shop(
                shop_ids[0], date_from_ms, date_to_ms, unit_ms=unit_ms
            )
        try:
            return self._fetch_expenses_bulk(
                shop_ids, date_from_ms, date_to_ms, unit_ms=unit_ms
            )
        except ShopUnavailableError:
            self.warnings.append(
                "Сводная выгрузка недоступна, загружаем магазины по одному (медленнее)."
            )
            return self._collect_per_shop(
                shop_ids,
                names,
                lambda sid: self._fetch_expenses_for_shop(
                    sid, date_from_ms, date_to_ms, unit_ms=unit_ms
                ),
            )

    def _fetch_dated_records(
        self,
        available: list[int],
        names: dict[int, str],
        date_from_ms: Optional[int],
        date_to_ms: Optional[int],
        fetcher: Callable[..., list[dict[str, Any]]],
        filter_fn: Callable[[list[dict[str, Any]]], list[dict[str, Any]]],
        label: str,
    ) -> list[dict[str, Any]]:
        items = fetcher(available, names, date_from_ms, date_to_ms, unit_ms=True)
        if items or (date_from_ms is None and date_to_ms is None):
            return items

        items = fetcher(available, names, date_from_ms, date_to_ms, unit_ms=False)
        if items:
            self.warnings.append(
                f"{label}: API принял границы периода в секундах (автоподбор формата)."
            )
            return items

        raw = fetcher(available, names, None, None, unit_ms=True)
        if not raw:
            return []

        filtered = filter_fn(raw)
        if filtered:
            self.warnings.append(
                f"{label}: серверный фильтр дат не вернул строки; "
                "применена локальная фильтрация по дате записи."
            )
            return filtered

        self.warnings.append(
            f"{label}: в API найдено {len(raw)} записей, но ни одна не попадает "
            "в выбранный период. Расширьте диапазон дат."
        )
        return []

    def _filter_available_shops(
        self,
        shop_ids: list[int],
        shop_names: dict[int, str],
        probe: Callable[[int], bool],
    ) -> list[int]:
        available: list[int] = []
        unavailable: list[str] = []
        for shop_id in shop_ids:
            label = shop_names.get(shop_id) or f"ID {shop_id}"
            if probe(shop_id):
                available.append(shop_id)
            else:
                unavailable.append(label)
            if PAUSE_BETWEEN_SHOPS_SEC > 0:
                time.sleep(PAUSE_BETWEEN_SHOPS_SEC)
        if unavailable:
            self.warnings.append(
                "Пропущены недоступные магазины (нет доступа в финансовом API): "
                + ", ".join(unavailable)
            )
        return available

    def _collect_per_shop(
        self,
        shop_ids: list[int],
        shop_names: dict[int, str],
        fetcher: Callable[[int], list[dict[str, Any]]],
    ) -> list[dict[str, Any]]:
        combined: list[dict[str, Any]] = []
        unavailable: list[str] = []
        for shop_id in shop_ids:
            label = shop_names.get(shop_id) or f"ID {shop_id}"
            try:
                combined.extend(fetcher(shop_id))
            except ShopUnavailableError:
                unavailable.append(label)
                continue
            finally:
                if PAUSE_BETWEEN_SHOPS_SEC > 0:
                    time.sleep(PAUSE_BETWEEN_SHOPS_SEC)
        if unavailable:
            self.warnings.append(
                "Пропущены недоступные магазины (нет доступа в финансовом API): "
                + ", ".join(unavailable)
            )
        return combined

    def fetch_finance_orders(
        self,
        date_from_ms: Optional[int],
        date_to_ms: Optional[int],
        shop_ids: list[int],
        shop_names: Optional[dict[int, str]] = None,
    ) -> list[dict[str, Any]]:
        if not shop_ids:
            raise RuntimeError("Для отчёта по продажам нужен хотя бы один магазин (shopIds).")

        names = shop_names or {sid: f"ID {sid}" for sid in shop_ids}
        available = self._filter_available_shops(
            shop_ids,
            names,
            self._probe_shop_for_orders,
        )
        if not available:
            raise RuntimeError(
                "Нет доступных магазинов для отчёта по продажам. "
                + (self.warnings[-1] if self.warnings else "")
            )

        return self._fetch_dated_records(
            available,
            names,
            date_from_ms,
            date_to_ms,
            self._fetch_orders_for_shops,
            lambda rows: _filter_by_date_range(
                rows, date_from_ms, date_to_ms, _order_item_timestamp_ms
            ),
            "Продажи",
        )

    def fetch_finance_expenses(
        self,
        date_from_ms: Optional[int],
        date_to_ms: Optional[int],
        shop_ids: list[int],
        shop_names: Optional[dict[int, str]] = None,
    ) -> list[dict[str, Any]]:
        if not shop_ids:
            raise RuntimeError("Для отчёта по услугам нужен хотя бы один магазин.")

        names = shop_names or {sid: f"ID {sid}" for sid in shop_ids}
        available = self._filter_available_shops(
            shop_ids,
            names,
            self._probe_shop_for_expenses,
        )
        if not available:
            raise RuntimeError(
                "Нет доступных магазинов для отчёта по услугам. "
                + (self.warnings[-1] if self.warnings else "")
            )

        return self._fetch_dated_records(
            available,
            names,
            date_from_ms,
            date_to_ms,
            self._fetch_expenses_for_shops,
            lambda rows: _filter_by_date_range(
                rows, date_from_ms, date_to_ms, _payment_timestamp_ms
            ),
            "Услуги",
        )

    def fetch_sku_stocks(self) -> list[dict[str, Any]]:
        data = self.get("/v2/fbs/sku/stocks")
        payload = (data or {}).get("payload") or {}
        return payload.get("skuAmountList") or []

    def fetch_fbo_invoices(self, shop_ids: list[int]) -> list[dict[str, Any]]:
        """FBO-накладные с составом (для «К отправке» и себестоимости)."""
        invoices: list[dict[str, Any]] = []
        page = 0
        while page < MAX_PAGES:
            data = self.get("/v1/invoice", {"page": page, "size": 50})
            batch: list[dict[str, Any]] = []
            if isinstance(data, list):
                batch = [x for x in data if isinstance(x, dict)]
            elif isinstance(data, dict):
                payload = data.get("payload")
                if isinstance(payload, list):
                    batch = [x for x in payload if isinstance(x, dict)]
            if not batch:
                break
            invoices.extend(batch)
            if len(batch) < 50:
                break
            page += 1

        if invoices:
            return self._hydrate_fbo_invoice_products(invoices)

        combined: list[dict[str, Any]] = []
        for shop_id in shop_ids:
            shop_page = 0
            while shop_page < MAX_PAGES:
                try:
                    data = self.get(
                        f"/v1/shop/{shop_id}/invoice",
                        {"page": shop_page, "size": 50},
                    )
                except RuntimeError:
                    break
                batch = data if isinstance(data, list) else []
                batch = [x for x in batch if isinstance(x, dict)]
                if not batch:
                    break
                for inv in batch:
                    inv = dict(inv)
                    inv.setdefault("shopId", shop_id)
                    combined.append(inv)
                if len(batch) < 50:
                    break
                shop_page += 1
            if PAUSE_BETWEEN_SHOPS_SEC > 0:
                time.sleep(PAUSE_BETWEEN_SHOPS_SEC)
        return self._hydrate_fbo_invoice_products(combined)

    def _hydrate_fbo_invoice_products(
        self,
        invoices: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        hydrated: list[dict[str, Any]] = []
        for invoice in invoices:
            inv = dict(invoice)
            products = inv.get("productForInvoiceDto")
            if isinstance(products, list) and products:
                hydrated.append(inv)
                continue
            if not _invoice_is_pending(inv):
                hydrated.append(inv)
                continue
            shop_raw = inv.get("shopId")
            inv_id = inv.get("id")
            if shop_raw is None or inv_id is None:
                hydrated.append(inv)
                continue
            try:
                shop_id = int(shop_raw)
                invoice_id = int(inv_id)
            except (TypeError, ValueError):
                hydrated.append(inv)
                continue
            try:
                data = self.get(
                    f"/v1/shop/{shop_id}/invoice/products",
                    {"invoiceId": invoice_id},
                )
                if isinstance(data, list):
                    inv["productForInvoiceDto"] = data
            except RuntimeError:
                pass
            hydrated.append(inv)
        return hydrated

    def fetch_product_catalog_index(self, shop_ids: list[int]) -> ProductCatalogIndex:
        index = ProductCatalogIndex.empty()
        if not shop_ids:
            return index
        for shop_id in shop_ids:
            page = 0
            loaded = 0
            while page < MAX_PAGES:
                data = self.get(
                    f"/v1/product/shop/{shop_id}",
                    {
                        "page": page,
                        "size": PRODUCT_PAGE_SIZE,
                        "filter": "ALL",
                    },
                )
                products = (data or {}).get("productList") or []
                if not products:
                    break
                index.add_shop_products(shop_id, products)
                loaded += len(products)
                total = (data or {}).get("totalProductsAmount")
                if total is not None and loaded >= int(total):
                    break
                if len(products) < PRODUCT_PAGE_SIZE:
                    break
                page += 1
            if PAUSE_BETWEEN_SHOPS_SEC > 0:
                time.sleep(PAUSE_BETWEEN_SHOPS_SEC)
        if index.size:
            self.warnings.append(
                "Часть колонок дополнена из каталога товаров Uzum (/v1/product/shop)."
            )
        return index

    def _accumulate_item_qty(
        self,
        item: dict[str, Any],
        catalog: ProductCatalogIndex,
        totals: dict[int, int],
    ) -> None:
        cat = catalog.lookup_order_item(item)
        sku_id: Optional[int] = None
        if cat and cat.sku_id is not None:
            sku_id = cat.sku_id
        else:
            raw = item.get("skuId")
            if raw is not None:
                try:
                    sku_id = int(raw)
                except (TypeError, ValueError):
                    sku_id = None
        if sku_id is None:
            return
        try:
            qty = int(item.get("amount") or 1)
        except (TypeError, ValueError):
            qty = 1
        if qty < 0:
            qty = 0
        totals[sku_id] = totals.get(sku_id, 0) + qty

    def _fetch_scheme_order_qty_totals_for_shop(
        self,
        shop_id: int,
        date_from_ms: int,
        date_to_ms: int,
        scheme: str,
        catalog: ProductCatalogIndex,
        *,
        unit_ms: bool = True,
    ) -> dict[int, int]:
        totals: dict[int, int] = {}
        page = 0
        while page < MAX_PAGES:
            params: dict[str, Any] = {
                "shopIds": [shop_id],
                "page": page,
                "size": FBS_ORDERS_PAGE_SIZE,
                "status": "COMPLETED",
                "scheme": scheme,
            }
            _apply_date_query_params(params, date_from_ms, date_to_ms, unit_ms=unit_ms)
            data = self.get("/v2/fbs/orders", params)
            payload = (data or {}).get("payload") or {}
            orders = payload.get("orders") or []
            if not orders:
                break
            for order in orders:
                if not isinstance(order, dict):
                    continue
                for item in order.get("orderItems") or []:
                    if isinstance(item, dict):
                        self._accumulate_item_qty(item, catalog, totals)
            total_amount = payload.get("totalAmount")
            if total_amount is not None and (page + 1) * FBS_ORDERS_PAGE_SIZE >= int(
                total_amount
            ):
                break
            if len(orders) < FBS_ORDERS_PAGE_SIZE:
                break
            page += 1
        return totals

    def _fetch_scheme_order_qty_totals(
        self,
        shop_ids: list[int],
        date_from_ms: int,
        date_to_ms: int,
        scheme: str,
        catalog: ProductCatalogIndex,
    ) -> dict[int, int]:
        totals: dict[int, int] = {}
        for shop_id in shop_ids:
            shop_totals: Optional[dict[int, int]] = None
            last_error: Optional[RuntimeError] = None
            for unit_ms in (True, False):
                try:
                    shop_totals = self._fetch_scheme_order_qty_totals_for_shop(
                        shop_id,
                        date_from_ms,
                        date_to_ms,
                        scheme,
                        catalog,
                        unit_ms=unit_ms,
                    )
                    break
                except RuntimeError as exc:
                    last_error = exc
                    if "400" not in str(exc):
                        raise
            if shop_totals is None:
                if last_error:
                    raise last_error
                shop_totals = {}
            for sku_id, qty in shop_totals.items():
                totals[sku_id] = totals.get(sku_id, 0) + qty
            if PAUSE_BETWEEN_SHOPS_SEC > 0:
                time.sleep(PAUSE_BETWEEN_SHOPS_SEC)
        return totals

    @staticmethod
    def _catalog_avg_daily_sales_by_sku(catalog: ProductCatalogIndex) -> dict[int, float]:
        out: dict[int, float] = {}
        for entry in catalog.by_sku_id_global.values():
            sku_id = entry.sku_id
            if sku_id is None or sku_id in out:
                continue
            avg = _safe_float(entry.avg_daily_sales)
            if avg is not None and avg > 0:
                out[sku_id] = avg
        return out

    def fetch_avg_daily_sales_15d_by_sku(
        self,
        shop_ids: list[int],
        catalog: ProductCatalogIndex,
        shop_names: dict[int, str],
    ) -> dict[int, float]:
        """Среднесуточные продажи за 15 дней по skuId (приоритет — заказы FBO)."""
        date_from_ms, date_to_ms = last_n_days_range_ms(SALES_LOOKBACK_DAYS)
        totals: dict[int, int] = {}
        try:
            totals = self._fetch_scheme_order_qty_totals(
                shop_ids, date_from_ms, date_to_ms, "FBO", catalog
            )
        except RuntimeError as exc:
            logger.warning("FBS orders for 15d sales failed: %s", exc)
            self.warnings.append(
                f"Заказы FBO за {SALES_LOOKBACK_DAYS} дн. недоступны ({exc}). "
                "Среднесуточные продажи взяты из каталога или финансового API."
            )
        if totals:
            return {
                sku_id: total / SALES_LOOKBACK_DAYS
                for sku_id, total in totals.items()
                if total > 0
            }

        catalog_rates = self._catalog_avg_daily_sales_by_sku(catalog)
        if catalog_rates:
            self.warnings.append(
                f"«Среднесуточные продажи FBO за {SALES_LOOKBACK_DAYS} дней»: "
                "использовано поле avgdsales из каталога товаров."
            )
            return catalog_rates

        for item in self.fetch_finance_orders(
            date_from_ms, date_to_ms, shop_ids, shop_names
        ):
            self._accumulate_item_qty(item, catalog, totals)
        if totals:
            self.warnings.append(
                f"«Среднесуточные продажи FBO за {SALES_LOOKBACK_DAYS} дней»: "
                "использованы финансовые продажи за период."
            )
            return {
                sku_id: total / SALES_LOOKBACK_DAYS
                for sku_id, total in totals.items()
                if total > 0
            }
        return {}


def _sales_sku_label(item: dict[str, Any], cat: Optional[SkuCatalogEntry]) -> str:
    """Human-readable seller SKU (e.g. SKLPTR-SLIP-БЕЛЫЙ-L), not numeric skuId."""
    if cat:
        for candidate in (
            cat.article,
            cat.seller_item_code,
            cat.sku_full_title,
            cat.sku_title,
        ):
            if candidate:
                return candidate
    for key in ("skuCharValue", "skuCharTitle", "skuTitle"):
        value = item.get(key)
        if value:
            return str(value)
    if cat and cat.sku_title:
        return cat.sku_title
    return ""


def _sales_money_fields(item: dict[str, Any]) -> tuple[Any, Any, Any, Any, Any]:
    # Цена — sellerPrice/sellPrice; выручка 0 при возврате; net = выручка − комиссия − логистика.
    price = item.get("sellerPrice")
    if price is None:
        price = item.get("sellPrice")

    returns_raw = item.get("amountReturns")
    try:
        has_return = int(returns_raw or 0) > 0
    except (TypeError, ValueError):
        has_return = False

    revenue: Any = 0 if has_return else price

    def _as_int(value: Any) -> int:
        if value is None or value == "":
            return 0
        try:
            return int(value)
        except (TypeError, ValueError):
            return 0

    revenue_net = _as_int(revenue) - _as_int(item.get("commission")) - _as_int(
        item.get("logisticDeliveryFee")
    )
    return price, revenue, revenue_net, item.get("commission"), item.get("logisticDeliveryFee")


def _expense_source_label(raw: Any, service_name: str = "") -> str:
    """Источник услуги в формате кабинета Uzum: «Склад» / «Маркетинг»."""
    if raw in (None, ""):
        text = ""
    else:
        text = str(raw).strip()
    if text:
        mapped = EXPENSE_SOURCE_RU.get(text.upper())
        if mapped:
            return mapped
        lower = text.lower().replace("\u2019", "'").replace("\u02bb", "'")
        if lower in ("склад", "sklad", "ombor", "storage", "warehouse"):
            return "Склад"
        if lower in ("маркетинг", "marketing", "reklama", "advertising"):
            return "Маркетинг"
    svc = str(service_name or "").upper()
    if any(token in svc for token in ("РЕКЛАМ", "MARKET", "BID", "PROMO", "ADS")):
        return "Маркетинг"
    if any(token in svc for token in ("ХРАНЕН", "STORAGE", "СКЛАД", "WAREHOUSE")):
        return "Склад"
    return text


def _build_sales_rows(
    items: list[dict[str, Any]],
    shop_names: dict[int, str],
    catalog: Optional[ProductCatalogIndex] = None,
) -> tuple[list[ExportColumn], list[dict[str, Any]]]:
    unmapped = {"Промокод (сумы)"}
    columns = [ExportColumn(name, name not in unmapped) for name in SALES_COLUMN_ORDER]
    assert {c.name for c in columns} == CANONICAL_SALES

    rows: list[dict[str, Any]] = []
    for item in items:
        shop_id = item.get("shopId")
        cat = catalog.lookup_order_item(item) if catalog else None
        price, revenue, revenue_net, commission, logistics = _sales_money_fields(item)
        rows.append(
            {
                "Статус": SALES_STATUS_RU.get(str(item.get("status") or ""), item.get("status") or ""),
                "Дата создания": format_datetime(item.get("date")),
                "Дата получения": format_datetime(item.get("dateIssued")),
                "№ заказа": item.get("orderId"),
                "Штрихкод": (cat.barcode if cat else "") or _format_barcode(item.get("barcode")),
                "SKU": _sales_sku_label(item, cat),
                "Наименование": item.get("productTitle") or item.get("skuTitle") or (cat.product_title if cat else ""),
                "Категория": (cat.category if cat else ""),
                "Количество": item.get("amount"),
                "Возвраты": item.get("amountReturns"),
                "Выручка (сумы)": revenue,
                "Выручка с вычетом комиссии и логистики (сумы)": revenue_net,
                "Комиссия маркетплейса (сумы)": commission,
                "Цена (сумы)": price,
                "Промокод (сумы)": "",
                "Себестоимость (сумы)": item.get("purchasePrice")
                if item.get("purchasePrice") is not None
                else (cat.purchase_price if cat else None),
                "Логистический сбор": item.get("logisticDeliveryFee"),
                "_shop": _shop_label(shop_names, shop_id),
            }
        )
    return columns, rows


def _build_expenses_rows(payments: list[dict[str, Any]], shop_names: dict[int, str]) -> tuple[list[ExportColumn], list[dict[str, Any]]]:
    columns = [ExportColumn(name, True) for name in EXPENSES_COLUMN_ORDER]
    assert {c.name for c in columns} == CANONICAL_EXPENSES

    rows: list[dict[str, Any]] = []
    for p in payments:
        shop_id = p.get("shopId")
        service_name = str(p.get("name") or "")
        rows.append(
            {
                "Источник": _expense_source_label(p.get("source"), service_name),
                "Услуга": service_name,
                "Статус": EXPENSE_STATUS_RU.get(str(p.get("status") or ""), p.get("status") or ""),
                "ID операции": p.get("id") if p.get("id") is not None else p.get("externalId"),
                "Дата списания": format_datetime(p.get("dateService") or p.get("dateCreated")),
                "Стоимость (сумы)": p.get("paymentPrice"),
                "Количество": p.get("amount"),
                "Сумма (сумы)": p.get("paymentPrice"),
                "Тип операции": EXPENSE_TYPE_RU.get(str(p.get("type") or ""), p.get("type") or ""),
                "_shop": _shop_label(shop_names, shop_id),
            }
        )
    return columns, rows


def _storage_row_dict(
    cat: SkuCatalogEntry,
    shop_names: dict[int, str],
    sales_15d_by_sku: Optional[dict[int, float]],
    fbo_stock_override: Any = None,
) -> dict[str, Any]:
    sku_id = cat.sku_id
    avg_sales_15d: Any = ""
    if sku_id is not None and sales_15d_by_sku:
        rate = sales_15d_by_sku.get(sku_id)
        if rate is not None:
            avg_sales_15d = _format_sales_rate(rate)
    if avg_sales_15d == "" and cat.avg_daily_sales not in (None, ""):
        avg_sales_15d = _format_sales_rate(float(cat.avg_daily_sales))

    fbo_stock = fbo_stock_override
    if fbo_stock in (None, ""):
        fbo_stock = cat.quantity_active

    avg_stock_15d = cat.avg_daily_stock
    if avg_stock_15d in (None, ""):
        avg_stock_15d = fbo_stock

    turnover = _storage_turnover_days(avg_stock_15d, avg_sales_15d)

    return {
        "Магазин": _shop_label(shop_names, cat.shop_id),
        "Название товара": cat.product_title,
        "ID товара": cat.product_id,
        "SKU": _sales_sku_label({}, cat),
        "Штрихкод": cat.barcode,
        "Габаритная группа": cat.dimensional_group,
        "Среднесуточные остатки FBO за 15 дней, шт": avg_stock_15d,
        "Среднесуточные продажи FBO за 15 дней, шт": avg_sales_15d,
        "Оборачиваемость, дней": turnover,
        "Хранение": _storage_type_label(cat),
        "За хранение 1 единицы 1 день, сум": _storage_price_per_unit(cat),
        "Остатки FBO (всего в продаже и на СДХ), шт": fbo_stock,
        "Всего за хранение 1 день, сум": _storage_fee_1_day(cat, fbo_stock),
        "Всего за хранение последние 30 дней, сум": _storage_fee_30d(cat, fbo_stock),
    }


def _build_storage_rows(
    stocks: list[dict[str, Any]],
    shop_names: dict[int, str],
    catalog: Optional[ProductCatalogIndex] = None,
    sales_15d_by_sku: Optional[dict[int, float]] = None,
    shop_ids: Optional[list[int]] = None,
) -> tuple[list[ExportColumn], list[dict[str, Any]]]:
    columns = [ExportColumn(name, True) for name in STORAGE_COLUMN_ORDER]
    assert {c.name for c in columns} == CANONICAL_STORAGE

    allowed_shops = set(shop_ids) if shop_ids else None
    rows: list[dict[str, Any]] = []
    seen: set[tuple[int, int]] = set()

    if catalog:
        catalog.enrich_storage_fields()
        for (shop_id, sku_id), entry in catalog.by_sku_id.items():
            if allowed_shops is not None and shop_id not in allowed_shops:
                continue
            key = (shop_id, sku_id)
            if key in seen:
                continue
            if not _include_in_storage_report(entry):
                continue
            seen.add(key)
            rows.append(_storage_row_dict(entry, shop_names, sales_15d_by_sku))

    for s in stocks:
        cat = catalog.lookup_stock_row(s) if catalog else None
        if cat is None or cat.sku_id is None:
            continue
        key = (cat.shop_id, cat.sku_id)
        if key in seen:
            continue
        if allowed_shops is not None and cat.shop_id not in allowed_shops:
            continue
        seen.add(key)
        fbo_stock = s.get("amount")
        if cat.quantity_active is not None:
            fbo_stock = cat.quantity_active
        rows.append(
            _storage_row_dict(cat, shop_names, sales_15d_by_sku, fbo_stock_override=fbo_stock)
        )

    return columns, rows


def _leftout_dimensional_group(label: str) -> str:
    if label:
        return label
    return "Неопределенная"


def _leftout_storage_label(cat: SkuCatalogEntry) -> str:
    if cat.pstorage is True:
        return "Платное"
    if cat.pstorage is False:
        return "Бесплатное"
    return "Нет данных"


def _leftout_status_label(cat: SkuCatalogEntry) -> str:
    in_sale = _qty_int(cat.quantity_active)
    returned = _qty_int(cat.quantity_returned)
    if returned > 0 and in_sale == 0:
        return "Возврат"
    if in_sale > 0:
        return "В продаже"
    if cat.product_status:
        return cat.product_status
    return "Не в продаже"


def _leftout_storage_daily_fee(cat: SkuCatalogEntry, in_sale: int) -> Any:
    if cat.pstorage is not True:
        return 0
    price = _safe_float(cat.paid_storage_price_item)
    if price is None:
        return 0
    return int(round(price * in_sale))


def _leftout_to_ship_qty(
    cat: SkuCatalogEntry,
    fbs_stock_row: Optional[dict[str, Any]] = None,
) -> int:
    """К отправке: pending → FBS в каталоге → остаток FBS из /v2/fbs/sku/stocks."""
    for val in (
        cat.quantity_pending,
        cat.quantity_fbs,
        cat.quantity_additional,
        fbs_stock_row.get("amount") if fbs_stock_row else None,
    ):
        qty = _qty_int(val)
        if qty > 0:
            return qty
    return 0


def _leftout_purchase_price(cat: SkuCatalogEntry) -> Any:
    if cat.purchase_price is not None:
        return cat.purchase_price
    return ""


def _leftout_row_dict(
    cat: SkuCatalogEntry,
    sales_15d_by_sku: Optional[dict[int, float]] = None,
    fbs_stock_row: Optional[dict[str, Any]] = None,
) -> dict[str, Any]:
    to_ship = _leftout_to_ship_qty(cat, fbs_stock_row)
    in_sale = _qty_int(cat.quantity_active)
    returned = _qty_int(cat.quantity_returned)
    defect = _qty_int(cat.quantity_defected)
    on_photo = _qty_int(cat.quantity_on_photo_studio)
    sdh = 0

    cost = _leftout_purchase_price(cat)
    price = cat.sell_price
    cost_num = _safe_float(cost) if cost not in (None, "") else None
    price_num = _safe_float(price)
    total_cost = int(round(cost_num * in_sale)) if cost_num is not None and in_sale else 0
    total_price = int(round(price_num * in_sale)) if price_num is not None and in_sale else 0

    avg_sales: Any = ""
    if cat.sku_id is not None and sales_15d_by_sku:
        rate = sales_15d_by_sku.get(cat.sku_id)
        if rate is not None:
            avg_sales = _format_sales_rate(rate)
    if avg_sales == "" and cat.avg_daily_sales not in (None, ""):
        avg_sales = _format_sales_rate(float(cat.avg_daily_sales))

    avg_stock = cat.avg_daily_stock
    if avg_stock in (None, ""):
        avg_stock = in_sale if in_sale else ""

    turnover = _storage_turnover_days(avg_stock, avg_sales)
    tariff = cat.paid_storage_price_item if cat.paid_storage_price_item is not None else 0
    seller_sku = _sales_sku_label({}, cat)

    return {
        "ID": cat.sku_id,
        "Наименование": cat.product_title or cat.sku_full_title or cat.sku_title,
        "Штрихкод": cat.barcode,
        "SKU": seller_sku,
        "ID товара": cat.product_id,
        "К отправке": to_ship,
        "В продаже": in_sale,
        "Возврат": returned,
        "Брак": defect,
        "Себест. (сумы)": cost,
        "Стоимость продажи (сумы)": price if price is not None else "",
        "Общий остаток": in_sale,
        "Общая сумма остатков (сумы)": total_price,
        "Себест. (сумма) (сумы)": total_cost,
        "Стоимость продажи (сумма) (сумы)": total_price,
        "Остаток на СДХ": sdh,
        "Остаток на фотостудии": on_photo,
        "Остаток на СДХ (сумма) (сумы)": 0,
        "Доступно к отправке": "",
        "Статус": _leftout_status_label(cat),
        "Габаритная группа": _leftout_dimensional_group(cat.dimensional_group),
        "Среднесуточные остатки": avg_stock,
        "Среднесуточные продажи": avg_sales,
        "Оборачиваемость": turnover,
        "Хранение": _leftout_storage_label(cat),
        "Тариф, сум": tariff,
        "Стоимость хранения 1 дня, сум": _leftout_storage_daily_fee(cat, in_sale),
        "Ссылка на товар": cat.preview_image,
    }


def _include_in_leftout_report(entry: SkuCatalogEntry) -> bool:
    if entry.sku_id is not None and entry.barcode:
        return True
    for field in (
        entry.quantity_active,
        entry.quantity_pending,
        entry.quantity_defected,
        entry.quantity_returned,
        entry.quantity_on_photo_studio,
    ):
        if (_safe_float(field) or 0) > 0:
            return True
    return False


def _build_inventory_old_rows(
    stocks: list[dict[str, Any]],
    _shop_names: dict[int, str],
    catalog: Optional[ProductCatalogIndex] = None,
    shop_ids: Optional[list[int]] = None,
    sales_15d_by_sku: Optional[dict[int, float]] = None,
) -> tuple[list[ExportColumn], list[dict[str, Any]]]:
    columns = [ExportColumn(name, True) for name in LEFTOUT_API_COLUMN_ORDER]
    assert {c.name for c in columns} == CANONICAL_LEFTOUT_API

    stocks_by_sku: dict[int, dict[str, Any]] = {}
    for s in stocks:
        raw = s.get("skuId")
        if raw is None:
            continue
        try:
            stocks_by_sku[int(raw)] = s
        except (TypeError, ValueError):
            continue

    allowed_shops = set(shop_ids) if shop_ids else None
    rows: list[dict[str, Any]] = []
    seen: set[tuple[int, int]] = set()

    if catalog:
        catalog.enrich_storage_fields()
        for (shop_id, sku_id), entry in catalog.by_sku_id.items():
            if allowed_shops is not None and shop_id not in allowed_shops:
                continue
            key = (shop_id, sku_id)
            if key in seen:
                continue
            if not _include_in_leftout_report(entry):
                continue
            seen.add(key)
            fbs_row = stocks_by_sku.get(sku_id) if sku_id is not None else None
            rows.append(_leftout_row_dict(entry, sales_15d_by_sku, fbs_row))

    for s in stocks:
        cat = catalog.lookup_stock_row(s) if catalog else None
        if cat is None or cat.sku_id is None:
            continue
        key = (cat.shop_id, cat.sku_id)
        if key in seen:
            continue
        if allowed_shops is not None and cat.shop_id not in allowed_shops:
            continue
        seen.add(key)
        rows.append(_leftout_row_dict(cat, sales_15d_by_sku, s))

    return columns, rows


ReportBuilder = Callable[[UzumApiClient, Optional[int], Optional[int], Optional[list[int]]], tuple[list[ExportColumn], list[dict[str, Any]]]]


def _shops_name_map(client: UzumApiClient) -> dict[int, str]:
    result: dict[int, str] = {}
    for s in client.list_shops():
        raw_id = s.get("id")
        if raw_id is None:
            continue
        try:
            result[int(raw_id)] = str(s.get("name") or "")
        except (TypeError, ValueError):
            continue
    return result


def _build_sales(
    client: UzumApiClient,
    date_from_ms: Optional[int],
    date_to_ms: Optional[int],
    shop_ids: Optional[list[int]],
) -> tuple[list[ExportColumn], list[dict[str, Any]]]:
    shops = _shops_name_map(client)
    resolved_ids = client.resolve_shop_ids(shop_ids)
    items = client.fetch_finance_orders(date_from_ms, date_to_ms, resolved_ids, shops)
    catalog = client.fetch_product_catalog_index(resolved_ids)
    return _build_sales_rows(items, shops, catalog)


def _build_expenses(
    client: UzumApiClient,
    date_from_ms: Optional[int],
    date_to_ms: Optional[int],
    shop_ids: Optional[list[int]],
) -> tuple[list[ExportColumn], list[dict[str, Any]]]:
    shops = _shops_name_map(client)
    resolved_ids = client.resolve_shop_ids(shop_ids)
    payments = client.fetch_finance_expenses(date_from_ms, date_to_ms, resolved_ids, shops)
    return _build_expenses_rows(payments, shops)


def _build_storage(
    client: UzumApiClient,
    _date_from_ms: Optional[int],
    _date_to_ms: Optional[int],
    shop_ids: Optional[list[int]],
) -> tuple[list[ExportColumn], list[dict[str, Any]]]:
    shops = _shops_name_map(client)
    resolved_ids = client.resolve_shop_ids(shop_ids)
    catalog = client.fetch_product_catalog_index(resolved_ids)
    sales_15d = client.fetch_avg_daily_sales_15d_by_sku(resolved_ids, catalog, shops)
    stocks = client.fetch_sku_stocks()
    return _build_storage_rows(stocks, shops, catalog, sales_15d, resolved_ids)


def _build_inventory_old(
    client: UzumApiClient,
    _date_from_ms: Optional[int],
    _date_to_ms: Optional[int],
    shop_ids: Optional[list[int]],
) -> tuple[list[ExportColumn], list[dict[str, Any]]]:
    shops = _shops_name_map(client)
    resolved_ids = client.resolve_shop_ids(shop_ids)
    catalog = client.fetch_product_catalog_index(resolved_ids)
    lookback_from, lookback_to = last_n_days_range_ms(365)
    order_items = client.fetch_finance_orders(lookback_from, lookback_to, resolved_ids, shops)
    filled_prices = catalog.enrich_purchase_prices_from_orders(order_items)
    catalog.enrich_purchase_prices_from_product_skus()
    if filled_prices:
        client.warnings.append(
            "Себест. (сумы): подставлена себестоимость из финансовых продаж за 365 дней."
        )
    try:
        invoices = client.fetch_fbo_invoices(resolved_ids)
        to_ship_filled, invoice_prices = catalog.enrich_from_fbo_invoices(invoices)
        if to_ship_filled:
            client.warnings.append(
                "К отправке: дополнено из открытых FBO-накладных (/v1/invoice)."
            )
        if invoice_prices:
            client.warnings.append(
                "Себест. (сумы): дополнено из FBO-накладных."
            )
    except RuntimeError as exc:
        logger.warning("FBO invoices skipped for left-out report: %s", exc)
    sales_15d = client.fetch_avg_daily_sales_15d_by_sku(resolved_ids, catalog, shops)
    stocks = client.fetch_sku_stocks()
    return _build_inventory_old_rows(stocks, shops, catalog, resolved_ids, sales_15d)


REPORT_BUILDERS: dict[str, ReportBuilder] = {
    "sales": _build_sales,
    "expenses": _build_expenses,
    "storage": _build_storage,
    "inventory_old": _build_inventory_old,
}


def date_to_epoch_ms(date_str: Optional[str], end_of_day: bool = False) -> Optional[int]:
    """Alias: calendar day boundaries in Uzbekistan (GMT+5)."""
    return calendar_date_to_epoch_ms(date_str, end_of_day=end_of_day)


def build_report(
    report_type: str,
    api_key: str,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    shop_ids: Optional[list[int]] = None,
) -> tuple[list[ExportColumn], list[dict[str, Any]], str, list[str]]:
    builder = REPORT_BUILDERS.get(report_type)
    if not builder:
        raise ValueError(f"Unknown report type: {report_type}")
    client = UzumApiClient(api_key)
    date_from_ms = date_to_epoch_ms(date_from, end_of_day=False)
    date_to_ms = date_to_epoch_ms(date_to, end_of_day=True)
    columns, rows = builder(client, date_from_ms, date_to_ms, shop_ids)
    return columns, rows, FILE_NAMES.get(report_type, f"{report_type}.xlsx"), list(client.warnings)


def build_xlsx_bytes(
    report_type: str,
    columns: list[ExportColumn],
    rows: list[dict[str, Any]],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAMES.get(report_type, "Отчет")[:31]

    bold_font = Font(bold=True)
    for col_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col.name)
        if not col.mapped:
            cell.font = bold_font

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col in enumerate(columns, start=1):
            value = row.get(col.name, "")
            if value is None:
                value = ""
            ws.cell(row=row_idx, column=col_idx, value=value)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _columns_for_report_type(report_type: str) -> list[ExportColumn]:
    if report_type == "sales":
        unmapped = {"Промокод (сумы)"}
        return [ExportColumn(n, n not in unmapped) for n in SALES_COLUMN_ORDER]
    if report_type == "expenses":
        return [ExportColumn(n, True) for n in EXPENSES_COLUMN_ORDER]
    if report_type == "storage":
        return [ExportColumn(n, True) for n in STORAGE_COLUMN_ORDER]
    if report_type == "inventory_old":
        return [ExportColumn(n, True) for n in LEFTOUT_API_COLUMN_ORDER]
    raise ValueError(f"Unknown report type: {report_type}")


def report_metadata() -> list[dict[str, Any]]:
    """Describe export templates for the UI."""
    hints = {
        "sales": "sells-report",
        "expenses": "expenses-report",
        "storage": "seller-storage-report",
        "inventory_old": "left-out-report",
    }
    meta = []
    tz = timezone_metadata()
    for report_type, sheet in SHEET_NAMES.items():
        cols = _columns_for_report_type(report_type)
        meta.append(
            {
                "id": report_type,
                "sheet_name": sheet,
                "file_name": FILE_NAMES[report_type],
                "hint": hints[report_type],
                "columns": [{"name": c.name, "mapped": c.mapped} for c in cols],
                "needs_date_range": report_type in ("sales", "expenses"),
                **tz,
            }
        )
    return meta
