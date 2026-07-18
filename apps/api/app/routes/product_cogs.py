"""Product unit COGS: Uzum LK source + Profiboard manual overrides."""
from __future__ import annotations

import io
import logging
import math
import re
from datetime import date, datetime, timedelta
from typing import Optional
from uuid import UUID

import pandas as pd
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.db import get_db, qname
from app.deps import require_admin
from app.routes.imports import table_exists
from app.schemas.product_cogs import (
    ProductCogsHistoryEntry,
    ProductCogsHistoryResponse,
    ProductCogsItem,
    ProductCogsListResponse,
    ProductCogsTemplateUploadResponse,
    ProductCogsUpsertRequest,
)
from app.utils.barcode import barcode_norm_sql
from app.utils.product_image import resolve_product_image_url
from app.utils.metrics import get_status_conditions
from app.utils.shop_filter import normalize_shop, storage_barcode_filter_sql
from app.utils.trial_shop import assert_trial_shop_filter_allowed

logger = logging.getLogger(__name__)
router = APIRouter()

_status = get_status_conditions()
_STATUS_REVENUE_SQL = f"(({_status['completed']}) OR ({_status['processing']}))"

COGS_TEMPLATE_HEADER_MAP = {
    "штрихкод": "barcode",
    "shtrixkod": "barcode",
    "barcode": "barcode",
    "наименование": "product_name",
    "tovar nomi": "product_name",
    "sku": "sku",
    "себестоимость (актуальная)": "actual_cogs",
    "tannarx (aktual)": "actual_cogs",
    "себестоимость": "actual_cogs",
    "tannarx": "actual_cogs",
    "actual_cogs": "actual_cogs",
    "дата начиная с": "effective_from",
    "boshlanish sanasi": "effective_from",
    "дата": "effective_from",
    "effective_from": "effective_from",
}

COGS_TEMPLATE_LOCALES: dict[str, dict[str, object]] = {
    "ru": {
        "sheet_name": "Себестоимость",
        "empty_columns": ["Штрихкод", "Себестоимость (актуальная)"],
        "products_columns": ["Штрихкод", "Наименование", "SKU", "Себестоимость (актуальная)"],
        "products_instruction": "Заполните только колонку Себестоимость (актуальная). Данные в других колонках не меняйте.",
        "empty_instruction": "Заполните колонки по товару Штрихкод и Себестоимость (актуальная).",
        "empty_filename": "sebestoimost_pustoy_shablon.xlsx",
        "products_filename": "sebestoimost_tovary_shablon.xlsx",
    },
    "uz": {
        "sheet_name": "Tannarx",
        "empty_columns": ["Shtrixkod", "Tannarx (aktual)"],
        "products_columns": ["Shtrixkod", "Tovar nomi", "SKU", "Tannarx (aktual)"],
        "products_instruction": "Faqat Tannarx (aktual) ustunini to'ldiring. Boshqa ustunlardagi ma'lumotlarni o'zgartirmang.",
        "empty_instruction": "Har bir tovar uchun Shtrixkod va Tannarx (aktual) ustunlarini to'ldiring.",
        "empty_filename": "tannarx_bosh_shablon.xlsx",
        "products_filename": "tannarx_tovarlar_shabloni.xlsx",
    },
}


def _resolve_cogs_template_lang(lang: Optional[str]) -> str:
    normalized = (lang or "").strip().lower()
    if normalized in ("uz", "uzbek", "o'z", "oz"):
        return "uz"
    return "ru"


def _cogs_template_locale(lang: Optional[str]) -> dict[str, object]:
    return COGS_TEMPLATE_LOCALES[_resolve_cogs_template_lang(lang)]


def _str_val(v) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def _safe_float(v) -> Optional[float]:
    if v is None:
        return None
    try:
        n = float(v)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(n):
        return None
    return n


_CURRENCY_TOKEN_RE = re.compile(
    r"(?i)\b(сум|sum|uzs|rub|rur|usd|eur|so[ʻʼ''`]?m)\b"
)
_CURRENCY_SYMBOL_RE = re.compile(r"[₽$€£¥]")
_SPACE_CHARS_RE = re.compile(r"[\s\u00a0\u202f\u2007\u2009\u200a\u200b]+")
_SCIENTIFIC_RE = re.compile(r"^[+-]?\d+(?:\.\d+)?[eE][+-]?\d+$")


def _strip_excel_noise(s: str) -> str:
    """Убрать пробелы/неразрывные пробелы и типичный мусор вставки из Excel."""
    return _SPACE_CHARS_RE.sub("", (s or "").strip())


def _parse_num(v) -> Optional[float]:
    """Число из ячейки: 12500 / 12 500 / 12 500,50 / 12,500.50 / 12500 сум / 12'500."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return _safe_float(v)
    s = str(v).strip()
    if not s or s in ("-", "—", "–"):
        return None
    s = (
        s.replace("\u00a0", " ")
        .replace("\u202f", " ")
        .replace("\u2007", " ")
        .replace("\u2009", " ")
        .replace("\u200a", " ")
        .replace("\u200b", "")
        .replace("'", "")
        .replace("’", "")
        .replace("`", "")
    )
    s = _CURRENCY_TOKEN_RE.sub("", s)
    s = _CURRENCY_SYMBOL_RE.sub("", s)
    s = _SPACE_CHARS_RE.sub("", s.strip())
    if not s or s in ("-", "—", "–"):
        return None
    if s.lower() in ("nan", "inf", "-inf", "infinity", "-infinity"):
        return None
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            # 12.500,50 (EU)
            s = s.replace(".", "").replace(",", ".")
        else:
            # 12,500.50 (US)
            s = s.replace(",", "")
    elif "," in s:
        parts = s.split(",")
        if len(parts) == 2 and parts[1].isdigit() and len(parts[1]) <= 2:
            s = s.replace(",", ".")
        else:
            s = s.replace(",", "")
    return _safe_float(s)


def _is_blank_cell(v) -> bool:
    if v is None:
        return True
    try:
        if pd.isna(v):
            return True
    except (TypeError, ValueError):
        pass
    if isinstance(v, str):
        s = _SPACE_CHARS_RE.sub("", v).strip().lower()
        return not s or s in ("-", "—", "–", "nan", "none", "null")
    return False


def _is_filled_template_cogs_cell(v) -> bool:
    """Пустые ячейки шаблона не импортируем; явный 0 (строка «0») — можно."""
    if _is_blank_cell(v):
        return False
    if isinstance(v, str):
        # «12500 сум», «12 500» и т.п. — заполненные; после очистки мусора должна остаться цифра.
        return _parse_num(v) is not None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            if pd.isna(v):
                return False
        except (TypeError, ValueError):
            pass
        # Пустые числовые ячейки Excel иногда читаются как 0.0.
        return float(v) != 0.0
    return True


def _normalize_barcode_norm(raw) -> str:
    """Штрихкод как текст: убрать пробелы, .0, научную запись Excel (1.0001E+12)."""
    if raw is None:
        return ""
    try:
        if pd.isna(raw):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(raw, bool):
        return ""
    if isinstance(raw, int):
        return str(raw)
    if isinstance(raw, float):
        if not math.isfinite(raw):
            return ""
        # Excel часто отдаёт штрихкод float'ом; целые — без дробной части.
        if raw == int(raw) and abs(raw) < 1e20:
            return str(int(raw))
        return _strip_excel_noise(f"{raw:.0f}" if abs(raw) >= 1e11 else str(raw))

    bn = _strip_excel_noise(str(raw))
    if not bn:
        return ""
    # 1.000102783428E+12 / 1,000102783428E+12
    sci = bn.replace(",", ".")
    if _SCIENTIFIC_RE.match(sci):
        try:
            n = float(sci)
            if math.isfinite(n) and n == int(n) and abs(n) < 1e20:
                return str(int(n))
        except (TypeError, ValueError, OverflowError):
            pass
    if bn.endswith(".0") and bn[:-2].replace("-", "", 1).isdigit():
        bn = bn[:-2]
    return bn


def _normalize_template_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _parse_date_cell(value: object) -> Optional[str]:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if hasattr(value, "date") and callable(value.date):
        try:
            return value.date().isoformat()
        except Exception:
            pass
    s = str(value).strip()
    if not s or s.lower() in ("nan", "none", "-", "—"):
        return None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        return s
    m = re.fullmatch(r"(\d{2})\.(\d{2})\.(\d{4})", s)
    if m:
        return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    try:
        parsed = pd.to_datetime(s, dayfirst=True, errors="coerce")
        if pd.isna(parsed):
            return None
        return parsed.date().isoformat()
    except Exception:
        return None


def _save_product_cogs(
    db: Session,
    user_id: UUID,
    barcode_norm: str,
    actual_cogs: float,
    effective_from: str,
    *,
    barcode: Optional[str] = None,
    sku: Optional[str] = None,
    product_name: Optional[str] = None,
    shop: Optional[str] = None,
) -> None:
    safe_cogs = _safe_float(actual_cogs)
    if safe_cogs is None:
        raise HTTPException(status_code=400, detail="Invalid actual_cogs value")
    _ensure_manual_product_cogs_schema(db)
    db.execute(
        text(f"""
            INSERT INTO {qname("manual_product_cogs")} (
                user_id, barcode_norm, barcode, sku, product_name, shop, cogs_sum, effective_from, updated_at
            ) VALUES (
                CAST(:user_id AS uuid), :barcode_norm, :barcode, :sku, :product_name, :shop,
                CAST(:cogs_sum AS numeric(18,2)), CAST(:effective_from AS date), now()
            )
            ON CONFLICT (user_id, barcode_norm) DO UPDATE SET
                barcode = EXCLUDED.barcode,
                sku = EXCLUDED.sku,
                product_name = EXCLUDED.product_name,
                shop = EXCLUDED.shop,
                cogs_sum = EXCLUDED.cogs_sum,
                effective_from = EXCLUDED.effective_from,
                updated_at = now()
        """),
        {
            "user_id": str(user_id),
            "barcode_norm": barcode_norm,
            "barcode": barcode,
            "sku": sku,
            "product_name": product_name,
            "shop": shop,
            "cogs_sum": safe_cogs,
            "effective_from": effective_from,
        },
    )
    _append_cogs_history(db, user_id, barcode_norm, safe_cogs, effective_from)


def _extract_product_image_url_from_data(data: dict) -> Optional[str]:
    raw = data.get("Ссылка на товар") or data.get("Mahsulot havolasi") or data.get("Tovar havolasi")
    return resolve_product_image_url(raw)


def _ensure_manual_product_cogs_schema(db: Session) -> None:
    """Таблица ручной себестоимости + колонка effective_from (дата начиная с)."""
    tbl = qname("manual_product_cogs")
    if not table_exists(db, tbl):
        db.execute(
            text(
                f"""
                CREATE TABLE IF NOT EXISTS {tbl} (
                    id uuid PRIMARY KEY DEFAULT gen_random_uuid(),
                    user_id uuid NOT NULL,
                    barcode_norm text NOT NULL,
                    barcode text,
                    sku text,
                    product_name text,
                    shop text,
                    cogs_sum numeric(18,2) NOT NULL,
                    effective_from date,
                    created_at timestamptz NOT NULL DEFAULT now(),
                    updated_at timestamptz NOT NULL DEFAULT now(),
                    UNIQUE (user_id, barcode_norm)
                )
                """
            )
        )
        db.execute(
            text(f"CREATE INDEX IF NOT EXISTS ix_manual_product_cogs_user ON {tbl} (user_id)")
        )
    else:
        db.execute(text(f"ALTER TABLE {tbl} ADD COLUMN IF NOT EXISTS effective_from date"))


def _ensure_manual_product_cogs_history_schema(db: Session) -> None:
    tbl = qname("manual_product_cogs_history")
    if not table_exists(db, tbl):
        db.execute(
            text(
                f"""
                CREATE TABLE IF NOT EXISTS {tbl} (
                    id uuid PRIMARY KEY DEFAULT gen_random_uuid(),
                    user_id uuid NOT NULL,
                    barcode_norm text NOT NULL,
                    cogs_sum numeric(18,2) NOT NULL,
                    effective_from date NOT NULL,
                    calculation_source text NOT NULL DEFAULT 'profiboard',
                    created_at timestamptz NOT NULL DEFAULT now()
                )
                """
            )
        )
        db.execute(
            text(
                f"""
                CREATE INDEX IF NOT EXISTS ix_manual_product_cogs_history_user_barcode
                ON {tbl} (user_id, barcode_norm, effective_from DESC, created_at DESC)
                """
            )
        )


def _append_cogs_history(
    db: Session,
    user_id: UUID,
    barcode_norm: str,
    cogs_sum: float,
    effective_from: str,
) -> None:
    _ensure_manual_product_cogs_history_schema(db)
    db.execute(
        text(f"""
            INSERT INTO {qname("manual_product_cogs_history")} (
                user_id, barcode_norm, cogs_sum, effective_from, calculation_source
            ) VALUES (
                CAST(:user_id AS uuid), :barcode_norm,
                CAST(:cogs_sum AS numeric(18,2)), CAST(:effective_from AS date), 'profiboard'
            )
        """),
        {
            "user_id": str(user_id),
            "barcode_norm": barcode_norm,
            "cogs_sum": cogs_sum,
            "effective_from": effective_from,
        },
    )


def _get_first_sale_date(db: Session, user_id: UUID, barcode_norm: str) -> Optional[str]:
    fs_barcode_norm_expr = (
        "CASE "
        "WHEN RIGHT(COALESCE(fs.barcode_norm, ''), 2) = '.0' "
        "THEN LEFT(COALESCE(fs.barcode_norm, ''), LENGTH(COALESCE(fs.barcode_norm, '')) - 2) "
        "ELSE COALESCE(fs.barcode_norm, '') END"
    )
    row = db.execute(
        text(f"""
            SELECT MIN(fs.date_created::date) AS first_sale_date
            FROM {qname("fact_sales")} fs
            WHERE fs.user_id = CAST(:user_id AS uuid)
              AND fs.barcode_norm IS NOT NULL
              AND {fs_barcode_norm_expr} = :barcode_norm
        """),
        {"user_id": str(user_id), "barcode_norm": barcode_norm},
    ).fetchone()
    if not row or not row[0]:
        return None
    val = row[0]
    return val.isoformat() if hasattr(val, "isoformat") else str(val)


def _get_last_sale_date(db: Session, user_id: UUID, barcode_norm: str) -> Optional[str]:
    fs_barcode_norm_expr = (
        "CASE "
        "WHEN RIGHT(COALESCE(fs.barcode_norm, ''), 2) = '.0' "
        "THEN LEFT(COALESCE(fs.barcode_norm, ''), LENGTH(COALESCE(fs.barcode_norm, '')) - 2) "
        "ELSE COALESCE(fs.barcode_norm, '') END"
    )
    row = db.execute(
        text(f"""
            SELECT MAX(fs.date_created::date) AS last_sale_date
            FROM {qname("fact_sales")} fs
            WHERE fs.user_id = CAST(:user_id AS uuid)
              AND fs.barcode_norm IS NOT NULL
              AND {fs_barcode_norm_expr} = :barcode_norm
        """),
        {"user_id": str(user_id), "barcode_norm": barcode_norm},
    ).fetchone()
    if not row or not row[0]:
        return None
    val = row[0]
    return val.isoformat() if hasattr(val, "isoformat") else str(val)


def _iso_day_before(iso_date: str) -> str:
    return (date.fromisoformat(iso_date) - timedelta(days=1)).isoformat()


def _get_lk_cogs_for_barcode(db: Session, user_id: UUID, barcode_norm: str) -> Optional[float]:
    batch_row = db.execute(
        text(f"""
            SELECT upload_batch_id
            FROM {qname("fact_leftout_old_snapshot")}
            WHERE user_id = CAST(:user_id AS uuid)
            ORDER BY loaded_at DESC NULLS LAST
            LIMIT 1
        """),
        {"user_id": str(user_id)},
    ).fetchone()
    if batch_row and batch_row[0]:
        stg_row = db.execute(
            text(f"""
                SELECT sl.data, sl.cost_raw
                FROM {qname("stg_leftout_old")} sl
                WHERE sl.user_id = CAST(:user_id AS uuid)
                  AND sl.upload_batch_id = CAST(:batch_id AS uuid)
                  AND NULLIF(trim(regexp_replace(COALESCE(sl.barcode_raw, ''), '\\s+', '', 'g')), '') = :barcode_norm
                LIMIT 1
            """),
            {
                "user_id": str(user_id),
                "batch_id": str(batch_row[0]),
                "barcode_norm": barcode_norm,
            },
        ).fetchone()
        if stg_row:
            data = stg_row[0] if isinstance(stg_row[0], dict) else {}
            lk_cogs = _parse_num(data.get("Себест. (сумы)")) or _parse_num(stg_row[1])
            if lk_cogs is not None:
                return lk_cogs

    fs_barcode_norm_expr = (
        "CASE "
        "WHEN RIGHT(COALESCE(fs.barcode_norm, ''), 2) = '.0' "
        "THEN LEFT(COALESCE(fs.barcode_norm, ''), LENGTH(COALESCE(fs.barcode_norm, '')) - 2) "
        "ELSE COALESCE(fs.barcode_norm, '') END"
    )
    sales_row = db.execute(
        text(f"""
            WITH last_dates AS (
                SELECT {fs_barcode_norm_expr} AS barcode_norm, MAX(fs.date_created::date) AS last_date
                FROM {qname("fact_sales")} fs
                WHERE fs.user_id = CAST(:user_id AS uuid) AND fs.barcode_norm IS NOT NULL
                  AND {_STATUS_REVENUE_SQL}
                GROUP BY {fs_barcode_norm_expr}
            )
            SELECT COALESCE(MAX(fs.cogs_sum), 0) AS unit_cogs
            FROM {qname("fact_sales")} fs
            JOIN last_dates ld
              ON ld.barcode_norm = {fs_barcode_norm_expr}
             AND fs.date_created::date = ld.last_date
            WHERE fs.user_id = CAST(:user_id AS uuid)
              AND {fs_barcode_norm_expr} = :barcode_norm
              AND {_STATUS_REVENUE_SQL}
        """),
        {"user_id": str(user_id), "barcode_norm": barcode_norm},
    ).fetchone()
    if sales_row and sales_row[0]:
        unit_cogs = _safe_float(sales_row[0])
        if unit_cogs is not None and unit_cogs > 0:
            return unit_cogs
    return None


def _load_profiboard_periods(db: Session, user_id: UUID, barcode_norm: str) -> list[tuple[str, float]]:
    """Периоды Profiboard: (effective_from, cogs), по возрастанию даты."""
    periods: list[tuple[str, float]] = []
    seen_dates: set[str] = set()

    tbl = qname("manual_product_cogs_history")
    if table_exists(db, tbl):
        _ensure_manual_product_cogs_history_schema(db)
        rows = db.execute(
            text(f"""
                SELECT DISTINCT ON (effective_from)
                    effective_from, cogs_sum
                FROM {tbl}
                WHERE user_id = CAST(:user_id AS uuid) AND barcode_norm = :barcode_norm
                ORDER BY effective_from ASC, created_at DESC
            """),
            {"user_id": str(user_id), "barcode_norm": barcode_norm},
        ).fetchall()
        for row in rows:
            eff = row[0]
            if eff is None:
                continue
            date_str = eff.isoformat() if hasattr(eff, "isoformat") else str(eff).strip()
            if date_str:
                cogs_val = _safe_float(row[1]) or 0.0
                periods.append((date_str, cogs_val))
                seen_dates.add(date_str)

    manual_tbl = qname("manual_product_cogs")
    if table_exists(db, manual_tbl):
        db.execute(text(f"ALTER TABLE {manual_tbl} ADD COLUMN IF NOT EXISTS effective_from date"))
        row = db.execute(
            text(f"""
                SELECT effective_from, cogs_sum
                FROM {manual_tbl}
                WHERE user_id = CAST(:user_id AS uuid) AND barcode_norm = :barcode_norm
                LIMIT 1
            """),
            {"user_id": str(user_id), "barcode_norm": barcode_norm},
        ).fetchone()
        if row and row[0]:
            date_str = (
                row[0].isoformat() if hasattr(row[0], "isoformat") else str(row[0]).strip()
            )
            if date_str and date_str not in seen_dates:
                cogs_val = _safe_float(row[1]) or 0.0
                periods.append((date_str, cogs_val))

    periods.sort(key=lambda x: x[0])
    return periods


def _sync_manual_product_cogs_from_history(db: Session, user_id: UUID, barcode_norm: str) -> None:
    """После изменения истории — актуальная запись = последний период Profiboard или удаление."""
    periods = _load_profiboard_periods(db, user_id, barcode_norm)
    manual_tbl = qname("manual_product_cogs")
    if not table_exists(db, manual_tbl):
        return
    uid = str(user_id)
    if not periods:
        db.execute(
            text(f"""
                DELETE FROM {manual_tbl}
                WHERE user_id = CAST(:user_id AS uuid) AND barcode_norm = :barcode_norm
            """),
            {"user_id": uid, "barcode_norm": barcode_norm},
        )
        return
    eff_date, cogs = periods[-1]
    db.execute(
        text(f"""
            UPDATE {manual_tbl}
            SET cogs_sum = CAST(:cogs_sum AS numeric(18,2)),
                effective_from = CAST(:effective_from AS date),
                updated_at = now()
            WHERE user_id = CAST(:user_id AS uuid) AND barcode_norm = :barcode_norm
        """),
        {
            "user_id": uid,
            "barcode_norm": barcode_norm,
            "cogs_sum": cogs,
            "effective_from": eff_date,
        },
    )


def _load_latest_profiboard_by_barcode(db: Session, user_id: UUID) -> dict[str, dict]:
    """Последняя запись Profiboard по каждому штрихкоду (max effective_from)."""
    from app.utils.profiboard_cogs import load_latest_profiboard_by_barcode

    return load_latest_profiboard_by_barcode(db, user_id)


def _current_year_start_iso() -> str:
    today = date.today()
    return date(today.year, 1, 1).isoformat()


def _build_cogs_timeline(
    lk_cogs: Optional[float],
    profiboard_periods: list[tuple[str, float]],
) -> list[ProductCogsHistoryEntry]:
    """Хронология себестоимости: Uzum с начала года, Profiboard с даты изменения; конец — сегодня."""
    year_start = _current_year_start_iso()
    today = date.today().isoformat()
    segments: list[tuple[str, float, str]] = []
    first_pb_date = profiboard_periods[0][0] if profiboard_periods else None

    if lk_cogs is not None:
        if not first_pb_date or first_pb_date > year_start:
            segments.append((year_start, lk_cogs, "uzum"))

    for eff_date, cogs in profiboard_periods:
        segments.append((eff_date, cogs, "profiboard"))

    if not segments:
        return []

    items: list[ProductCogsHistoryEntry] = []
    for i, (start, cogs, source) in enumerate(segments):
        if i + 1 < len(segments):
            period_to = _iso_day_before(segments[i + 1][0])
        else:
            period_to = today
        if period_to < start:
            period_to = start
        items.append(
            ProductCogsHistoryEntry(
                period_from=start,
                period_to=period_to,
                cogs=cogs,
                calculation_source=source,
                effective_from=start if source == "profiboard" else None,
                can_delete=source == "profiboard",
            )
        )
    return items


def _load_manual_cogs_by_barcode(db: Session, user_id: UUID) -> dict[str, dict]:
    manual_by_barcode: dict[str, dict] = {}
    tbl = qname("manual_product_cogs")
    if not table_exists(db, tbl):
        return manual_by_barcode
    db.execute(text(f"ALTER TABLE {tbl} ADD COLUMN IF NOT EXISTS effective_from date"))
    manual_query = text(f"""
        SELECT barcode_norm, cogs_sum, effective_from, barcode, sku, product_name, shop
        FROM {tbl}
        WHERE user_id = CAST(:user_id AS uuid)
    """)
    for row in db.execute(manual_query, {"user_id": str(user_id)}).fetchall():
        bn = (row[0] or "").strip()
        if bn:
            eff = row[2]
            effective_from_str = None
            if eff is not None:
                effective_from_str = (
                    eff.isoformat() if hasattr(eff, "isoformat") else str(eff).strip()
                )
            manual_by_barcode[bn] = {
                "cogs": _safe_float(row[1]) or 0.0,
                "effective_from": effective_from_str,
                "barcode": row[3],
                "sku": row[4],
                "product_name": row[5],
                "shop": row[6],
            }
    return manual_by_barcode


def _load_storage_shop_by_barcode(db: Session, user_id: UUID) -> dict[str, str]:
    """Магазин по штрихкоду из seller-storage; пустой shop = товар в архиве (как в таблице товаров)."""
    storage_query = text(f"""
        SELECT
            COALESCE(
                fss.barcode_norm,
                NULLIF(trim(regexp_replace(COALESCE(fss.barcode, ''), '\\s+', '', 'g')), '')
            ) AS barcode_norm,
            MAX(NULLIF(trim(fss.shop_raw), '')) AS shop
        FROM {qname("fact_storage_snapshot")} fss
        WHERE fss.user_id = CAST(:user_id AS uuid)
          AND (fss.barcode_norm IS NOT NULL OR fss.barcode IS NOT NULL)
        GROUP BY
            fss.user_id,
            COALESCE(
                fss.barcode_norm,
                NULLIF(trim(regexp_replace(COALESCE(fss.barcode, ''), '\\s+', '', 'g')), '')
            )
        HAVING COALESCE(MAX(NULLIF(trim(fss.shop_raw), '')), '') <> ''
    """)
    storage_by_barcode: dict[str, str] = {}
    for row in db.execute(storage_query, {"user_id": str(user_id)}).fetchall():
        bn = (row[0] or "").strip() if row[0] else ""
        shop_val = _str_val(row[1])
        if bn and shop_val:
            storage_by_barcode[bn] = shop_val
    return storage_by_barcode


@router.get("/product-cogs", response_model=ProductCogsListResponse)
async def list_product_cogs(
    user_id: UUID = Depends(require_admin),
    shop: Optional[str] = Query(default=None, description="Shop name filter"),
    db: Session = Depends(get_db),
):
    """Себестоимость по штрихкоду: ЛК Uzum + актуальная (Profiboard)."""
    try:
        assert_trial_shop_filter_allowed(db, user_id, shop, None)
        shop_norm = normalize_shop(shop)

        params_batch = {"user_id": str(user_id)}
        shop_filter_sql = ""
        if shop_norm:
            params_batch["shop_norm"] = shop_norm
            shop_filter_sql = "\n              " + storage_barcode_filter_sql(
                "sl", prefix_and=True, outer_barcode_norm_expr=barcode_norm_sql("sl.barcode_raw")
            )

        # Последний батч left-out-report_old (как в shipment / products-table)
        if not shop_norm:
            batch_query = text(f"""
                SELECT upload_batch_id
                FROM {qname("fact_leftout_old_snapshot")}
                WHERE user_id = CAST(:user_id AS uuid)
                ORDER BY loaded_at DESC NULLS LAST
                LIMIT 1
            """)
        else:
            batch_query = text(f"""
                SELECT sl.upload_batch_id FROM {qname("stg_leftout_old")} sl
                WHERE sl.user_id = CAST(:user_id AS uuid)
                {shop_filter_sql}
                ORDER BY sl.upload_batch_id DESC NULLS LAST
                LIMIT 1
            """)
        batch_row = db.execute(batch_query, params_batch).fetchone()
        if not batch_row or not batch_row[0]:
            return ProductCogsListResponse(items=[])

        batch_id = str(batch_row[0])
        params_stg = {"user_id": str(user_id), "batch_id": batch_id}
        if shop_norm:
            params_stg["shop_norm"] = shop_norm

        stg_query = text(f"""
            SELECT sl.data, sl.barcode_raw, sl.cost_raw, sl.price_raw,
                   NULLIF(trim(regexp_replace(COALESCE(sl.barcode_raw, ''), '\\s+', '', 'g')), '') AS barcode_norm
            FROM {qname("stg_leftout_old")} sl
            WHERE sl.user_id = CAST(:user_id AS uuid)
              AND sl.upload_batch_id = CAST(:batch_id AS uuid)
            {shop_filter_sql if shop_norm else ""}
        """)
        stg_rows = db.execute(stg_query, params_stg).fetchall()

        fs_barcode_norm_expr = (
            "CASE "
            "WHEN RIGHT(COALESCE(fs.barcode_norm, ''), 2) = '.0' "
            "THEN LEFT(COALESCE(fs.barcode_norm, ''), LENGTH(COALESCE(fs.barcode_norm, '')) - 2) "
            "ELSE COALESCE(fs.barcode_norm, '') END"
        )
        sales_cogs_query = text(f"""
            WITH last_dates AS (
                SELECT {fs_barcode_norm_expr} AS barcode_norm, MAX(fs.date_created::date) AS last_date
                FROM {qname("fact_sales")} fs
                WHERE fs.user_id = CAST(:user_id AS uuid) AND fs.barcode_norm IS NOT NULL
                  AND {_STATUS_REVENUE_SQL}
                GROUP BY {fs_barcode_norm_expr}
            ),
            last_rows AS (
                SELECT {fs_barcode_norm_expr} AS barcode_norm, fs.cogs_sum, fs.revenue_sum, fs.qty, fs.returns_qty
                FROM {qname("fact_sales")} fs
                JOIN last_dates ld ON ld.barcode_norm = {fs_barcode_norm_expr} AND fs.date_created::date = ld.last_date
                WHERE fs.user_id = CAST(:user_id AS uuid) AND {_STATUS_REVENUE_SQL}
            )
            SELECT barcode_norm,
                   COALESCE(MAX(cogs_sum), 0) AS unit_cogs,
                   CASE WHEN COALESCE(SUM(GREATEST(COALESCE(qty,0)-COALESCE(returns_qty,0),0)),0) > 0
                        THEN SUM(COALESCE(revenue_sum,0)) / NULLIF(SUM(GREATEST(COALESCE(qty,0)-COALESCE(returns_qty,0),0)),0)
                        ELSE 0 END AS unit_price
            FROM last_rows
            GROUP BY barcode_norm
        """)
        sales_by_barcode: dict[str, dict] = {}
        for row in db.execute(sales_cogs_query, {"user_id": str(user_id)}).fetchall():
            bn = (row[0] or "").strip() if row[0] else ""
            if bn:
                sales_by_barcode[bn] = {
                    "cogs": _safe_float(row[1]),
                    "price": _safe_float(row[2]),
                }

        first_sale_query = text(f"""
            SELECT {fs_barcode_norm_expr} AS barcode_norm, MIN(fs.date_created::date) AS first_sale_date
            FROM {qname("fact_sales")} fs
            WHERE fs.user_id = CAST(:user_id AS uuid) AND fs.barcode_norm IS NOT NULL
            GROUP BY {fs_barcode_norm_expr}
        """)
        first_sale_by_barcode: dict[str, str] = {}
        for row in db.execute(first_sale_query, {"user_id": str(user_id)}).fetchall():
            bn = (row[0] or "").strip() if row[0] else ""
            if bn and row[1]:
                first_sale_by_barcode[bn] = (
                    row[1].isoformat() if hasattr(row[1], "isoformat") else str(row[1])
                )

        latest_profiboard_by_barcode = _load_latest_profiboard_by_barcode(db, user_id)
        storage_shop_by_barcode = _load_storage_shop_by_barcode(db, user_id)

        items: list[ProductCogsItem] = []
        for row in stg_rows:
            data = row[0]
            if not isinstance(data, dict):
                data = {}
            barcode_raw = _str_val(row[1])
            cost_raw = _str_val(row[2]) if len(row) > 2 else None
            price_raw = _str_val(row[3]) if len(row) > 3 else None
            barcode_norm = (row[4] or "").strip() if len(row) > 4 and row[4] else ""
            if not barcode_norm and barcode_raw:
                barcode_norm = _normalize_barcode_norm(barcode_raw)
            if not barcode_norm:
                continue

            product_shop = storage_shop_by_barcode.get(barcode_norm)
            if not product_shop:
                continue
            if shop_norm and normalize_shop(product_shop) != shop_norm:
                continue

            product_name = _str_val(data.get("Наименование") or data.get("Название товара"))
            sku = _str_val(data.get("SKU"))
            barcode = _str_val(data.get("Штрихкод")) or barcode_raw

            price = _parse_num(data.get("Стоимость продажи (сумы)")) or _parse_num(price_raw)
            lk_cogs = _parse_num(data.get("Себест. (сумы)")) or _parse_num(cost_raw)

            sales_row = sales_by_barcode.get(barcode_norm)
            if sales_row:
                sales_price = sales_row.get("price")
                sales_cogs = sales_row.get("cogs")
                if (price is None or price <= 0) and sales_price is not None and sales_price > 0:
                    price = sales_price
                if (lk_cogs is None or lk_cogs <= 0) and sales_cogs is not None and sales_cogs > 0:
                    lk_cogs = sales_cogs

            pb_row = latest_profiboard_by_barcode.get(barcode_norm)
            if pb_row:
                actual_cogs = _safe_float(pb_row["cogs"])
                effective_from_str = pb_row["effective_from"]
                used_cogs = actual_cogs
                calculation_source = "profiboard"
            else:
                actual_cogs = None
                effective_from_str = None
                used_cogs = lk_cogs
                calculation_source = "uzum"

            unit_margin = None
            if price is not None and used_cogs is not None:
                margin = price - used_cogs
                unit_margin = margin if math.isfinite(margin) else None

            date_str = effective_from_str

            items.append(
                ProductCogsItem(
                    barcode=barcode,
                    barcode_norm=barcode_norm,
                    product_name=product_name,
                    sku=sku,
                    product_image_url=_extract_product_image_url_from_data(data),
                    price=price,
                    lk_cogs=lk_cogs,
                    actual_cogs=actual_cogs,
                    unit_margin=unit_margin,
                    calculation_source=calculation_source,
                    date=date_str,
                    effective_from=effective_from_str,
                    first_sale_date=first_sale_by_barcode.get(barcode_norm),
                )
            )

        items.sort(key=lambda x: (x.product_name or "", x.barcode_norm))
        return ProductCogsListResponse(items=items)
    except HTTPException:
        raise
    except Exception as e:
        logger.error("list_product_cogs failed: %s", e, exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@router.put("/product-cogs/{barcode_norm}", response_model=ProductCogsItem)
async def upsert_product_cogs(
    barcode_norm: str,
    body: ProductCogsUpsertRequest,
    user_id: UUID = Depends(require_admin),
    db: Session = Depends(get_db),
):
    """Сохранить актуальную себестоимость (ЛК Profiboard)."""
    bn = _normalize_barcode_norm(barcode_norm)
    if not bn:
        raise HTTPException(status_code=400, detail="barcode_norm is required")
    if not body.effective_from:
        raise HTTPException(status_code=400, detail="effective_from is required")

    _save_product_cogs(
        db,
        user_id,
        bn,
        body.actual_cogs,
        body.effective_from,
        barcode=body.barcode,
        sku=body.sku,
        product_name=body.product_name,
        shop=body.shop,
    )
    db.commit()

    resp = await list_product_cogs(user_id=user_id, shop=body.shop, db=db)
    for item in resp.items:
        if item.barcode_norm == bn:
            return item

    return ProductCogsItem(
        barcode=body.barcode,
        barcode_norm=bn,
        product_name=body.product_name,
        sku=body.sku,
        price=None,
        lk_cogs=None,
        actual_cogs=body.actual_cogs,
        unit_margin=None,
        calculation_source="profiboard",
        date=body.effective_from,
        effective_from=body.effective_from,
    )


@router.delete("/product-cogs/{barcode_norm}")
async def delete_product_cogs(
    barcode_norm: str,
    user_id: UUID = Depends(require_admin),
    db: Session = Depends(get_db),
):
    """Удалить актуальную себестоимость — вернуться к ЛК Uzum."""
    bn = _normalize_barcode_norm(barcode_norm)
    if not bn:
        raise HTTPException(status_code=400, detail="barcode_norm is required")

    result = db.execute(
        text(f"""
            DELETE FROM {qname("manual_product_cogs")}
            WHERE user_id = CAST(:user_id AS uuid) AND barcode_norm = :barcode_norm
        """),
        {"user_id": str(user_id), "barcode_norm": bn},
    )
    db.commit()
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail="Not found")
    return {"ok": True}


@router.get("/product-cogs/{barcode_norm}/history", response_model=ProductCogsHistoryResponse)
async def get_product_cogs_history(
    barcode_norm: str,
    user_id: UUID = Depends(require_admin),
    product_name: Optional[str] = Query(default=None),
    sku: Optional[str] = Query(default=None),
    lk_cogs: Optional[float] = Query(default=None, description="Себестоимость ЛК Uzum (fallback)"),
    db: Session = Depends(get_db),
):
    """Себестоимость по периодам: Uzum с начала года, Profiboard с даты изменения, конец — сегодня."""
    bn = _normalize_barcode_norm(barcode_norm)
    if not bn:
        raise HTTPException(status_code=400, detail="barcode_norm is required")

    first_sale_date = _get_first_sale_date(db, user_id, bn)
    lk_value = lk_cogs if lk_cogs is not None and lk_cogs > 0 else _get_lk_cogs_for_barcode(db, user_id, bn)
    profiboard_periods = _load_profiboard_periods(db, user_id, bn)
    items = _build_cogs_timeline(lk_value, profiboard_periods)

    return ProductCogsHistoryResponse(
        product_name=product_name,
        sku=sku,
        first_sale_date=first_sale_date,
        items=items,
    )


@router.delete("/product-cogs/{barcode_norm}/history", response_model=ProductCogsHistoryResponse)
async def delete_product_cogs_history_period(
    barcode_norm: str,
    effective_from: str = Query(..., description="Дата начала периода Profiboard (YYYY-MM-DD)"),
    user_id: UUID = Depends(require_admin),
    product_name: Optional[str] = Query(default=None),
    sku: Optional[str] = Query(default=None),
    lk_cogs: Optional[float] = Query(default=None),
    db: Session = Depends(get_db),
):
    """Удалить период Profiboard из истории. Период ЛК Uzum не удаляется — восстанавливается при удалении записей поверх."""
    bn = _normalize_barcode_norm(barcode_norm)
    if not bn:
        raise HTTPException(status_code=400, detail="barcode_norm is required")
    eff = (effective_from or "").strip()
    if len(eff) != 10 or eff[4] != "-" or eff[7] != "-":
        raise HTTPException(status_code=400, detail="effective_from must be YYYY-MM-DD")

    _ensure_manual_product_cogs_history_schema(db)
    tbl = qname("manual_product_cogs_history")
    result = db.execute(
        text(f"""
            DELETE FROM {tbl}
            WHERE user_id = CAST(:user_id AS uuid)
              AND barcode_norm = :barcode_norm
              AND effective_from = CAST(:effective_from AS date)
        """),
        {"user_id": str(user_id), "barcode_norm": bn, "effective_from": eff},
    )
    manual_tbl = qname("manual_product_cogs")
    manual_deleted = 0
    if table_exists(db, manual_tbl):
        manual_deleted = (
            db.execute(
                text(f"""
                    DELETE FROM {manual_tbl}
                    WHERE user_id = CAST(:user_id AS uuid)
                      AND barcode_norm = :barcode_norm
                      AND effective_from = CAST(:effective_from AS date)
                """),
                {"user_id": str(user_id), "barcode_norm": bn, "effective_from": eff},
            ).rowcount
            or 0
        )
    if result.rowcount == 0 and manual_deleted == 0:
        raise HTTPException(status_code=404, detail="History period not found")

    _sync_manual_product_cogs_from_history(db, user_id, bn)
    db.commit()

    return await get_product_cogs_history(
        barcode_norm=bn,
        user_id=user_id,
        product_name=product_name,
        sku=sku,
        lk_cogs=lk_cogs,
        db=db,
    )


def _template_has_barcode_header(df: pd.DataFrame) -> bool:
    for col in df.columns:
        key = _normalize_template_header(col)
        if COGS_TEMPLATE_HEADER_MAP.get(key) == "barcode":
            return True
    return False


def _parse_cogs_template_rows(df: pd.DataFrame) -> list[dict[str, object]]:
    if df is None or df.empty:
        return []
    col_map: dict[str, str] = {}
    for col in df.columns:
        key = _normalize_template_header(col)
        mapped = COGS_TEMPLATE_HEADER_MAP.get(key)
        if mapped:
            col_map[mapped] = str(col)
    if "barcode" not in col_map:
        raise HTTPException(status_code=400, detail="В файле нет колонки «Штрихкод»")
    rows: list[dict[str, object]] = []
    for _, series in df.iterrows():
        rows.append({field: series.get(col_name) for field, col_name in col_map.items()})
    return rows


def _autosize_excel_columns(
    worksheet,
    *,
    min_width: float = 10,
    max_width: float = 60,
    start_row: int = 1,
    max_column: Optional[int] = None,
) -> None:
    from openpyxl.utils import get_column_letter

    for col_idx, column_cells in enumerate(worksheet.columns, start=1):
        if max_column is not None and col_idx > max_column:
            break
        max_len = 0
        for cell in column_cells:
            if cell.row < start_row:
                continue
            value = cell.value
            if value is None:
                continue
            max_len = max(max_len, len(str(value)))
        if max_len > 0:
            worksheet.column_dimensions[get_column_letter(col_idx)].width = min(
                max(max_len + 2, min_width),
                max_width,
            )


def _instruction_merge_end_column(instruction: str, table_column_count: int) -> int:
    # Default Excel column width fits ~8 characters; merge extra columns without resizing them.
    chars_per_col = 8
    cols_for_text = math.ceil(len(instruction) / chars_per_col)
    return max(table_column_count, cols_for_text)


def _write_cogs_template_xlsx(
    rows: list[dict[str, object]],
    columns: list[str],
    *,
    instruction: Optional[str] = None,
    sheet_name: str = "Себестоимость",
) -> io.BytesIO:
    from openpyxl.styles import Alignment

    df = pd.DataFrame(rows, columns=columns)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        startrow = 1 if instruction else 0
        df.to_excel(writer, index=False, sheet_name=sheet_name, startrow=startrow)
        ws = writer.sheets[sheet_name]
        autosize_start_row = 2 if instruction else 1
        _autosize_excel_columns(ws, start_row=autosize_start_row, max_column=len(columns))
        if instruction:
            merge_end_col = _instruction_merge_end_column(instruction, len(columns))
            cell = ws["A1"]
            cell.value = instruction
            cell.alignment = Alignment(wrap_text=False, vertical="center")
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=merge_end_col)
    buf.seek(0)
    return buf


def _read_cogs_template_dataframe(content: bytes) -> pd.DataFrame:
    read_kwargs = {"engine": "openpyxl", "dtype": str, "keep_default_na": False}
    raw = io.BytesIO(content)
    df = pd.read_excel(raw, **read_kwargs)
    if not _template_has_barcode_header(df):
        raw.seek(0)
        df = pd.read_excel(raw, header=1, **read_kwargs)
    return df


@router.get("/product-cogs/template")
async def download_product_cogs_template(
    user_id: UUID = Depends(require_admin),
    shop: Optional[str] = Query(default=None, description="Shop name filter"),
    kind: str = Query(default="products", description="products | empty"),
    lang: Optional[str] = Query(default=None, description="Template language: ru | uz"),
    db: Session = Depends(get_db),
):
    """Скачать XLSX-шаблон: products — список товаров; empty — только заголовки."""
    template_kind = (kind or "products").strip().lower()
    if template_kind not in ("products", "empty"):
        raise HTTPException(status_code=400, detail="kind must be products or empty")

    locale_cfg = _cogs_template_locale(lang)
    sheet_name = str(locale_cfg["sheet_name"])

    if template_kind == "empty":
        rows: list[dict[str, object]] = []
        columns = list(locale_cfg["empty_columns"])
        filename = str(locale_cfg["empty_filename"])
        instruction = str(locale_cfg["empty_instruction"])
    else:
        resp = await list_product_cogs(user_id=user_id, shop=shop, db=db)
        columns = list(locale_cfg["products_columns"])
        rows = [
            {
                columns[0]: item.barcode or item.barcode_norm,
                columns[1]: item.product_name or "",
                columns[2]: item.sku or "",
                columns[3]: "",
            }
            for item in resp.items
        ]
        filename = str(locale_cfg["products_filename"])
        instruction = str(locale_cfg["products_instruction"])

    buf = _write_cogs_template_xlsx(
        rows,
        columns,
        instruction=instruction,
        sheet_name=sheet_name,
    )
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/product-cogs/template", response_model=ProductCogsTemplateUploadResponse)
async def upload_product_cogs_template(
    user_id: UUID = Depends(require_admin),
    shop: Optional[str] = Query(default=None, description="Shop name filter"),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """Загрузить XLSX-шаблон с актуальной себестоимостью."""
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Файл пустой")
    try:
        df = _read_cogs_template_dataframe(content)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Некорректный XLSX: {exc}") from exc

    parsed_rows = _parse_cogs_template_rows(df)
    catalog_resp = await list_product_cogs(user_id=user_id, shop=shop, db=db)
    catalog_by_barcode = {item.barcode_norm: item for item in catalog_resp.items}

    pending: list[tuple[str, float, str, Optional[str], Optional[str], Optional[str], int]] = []
    skipped = 0
    errors: list[str] = []

    for excel_row_idx, row in enumerate(parsed_rows, start=2):
        barcode_raw = row.get("barcode")
        bn = _normalize_barcode_norm(barcode_raw)
        actual_cogs_raw = row.get("actual_cogs")
        cogs = _parse_num(actual_cogs_raw) if _is_filled_template_cogs_cell(actual_cogs_raw) else None
        eff = _parse_date_cell(row.get("effective_from"))
        product_name = _str_val(row.get("product_name"))
        sku = _str_val(row.get("sku"))

        if not bn:
            if cogs is None and not eff:
                skipped += 1
                continue
            errors.append(f"Строка {excel_row_idx}: не указан штрихкод")
            continue
        if cogs is None:
            skipped += 1
            continue
        if not eff:
            eff = date(date.today().year, 1, 1).isoformat()
        if cogs < 0:
            errors.append(f"Строка {excel_row_idx}: некорректная себестоимость")
            continue

        catalog_item = catalog_by_barcode.get(bn)
        pending.append(
            (
                bn,
                cogs,
                eff,
                (catalog_item.barcode if catalog_item else bn) or bn,
                sku or (catalog_item.sku if catalog_item else None),
                product_name or (catalog_item.product_name if catalog_item else None),
                excel_row_idx,
            )
        )

    pending.sort(key=lambda item: (item[0], item[2]))

    imported = 0
    for bn, cogs, eff, barcode, sku_val, product_name, excel_row_idx in pending:
        try:
            _save_product_cogs(
                db,
                user_id,
                bn,
                cogs,
                eff,
                barcode=barcode,
                sku=sku_val,
                product_name=product_name,
                shop=shop,
            )
            imported += 1
        except Exception as exc:
            logger.warning("cogs template row %s failed: %s", excel_row_idx, exc)
            errors.append(f"Строка {excel_row_idx}: не удалось сохранить")

    if imported > 0:
        db.commit()
    else:
        db.rollback()

    return ProductCogsTemplateUploadResponse(
        imported=imported,
        skipped=skipped,
        errors=errors[:30],
    )
